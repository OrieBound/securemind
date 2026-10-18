"""Build the SecureMind architecture Word document.

Regenerate after edits:

    python scripts/build-architecture-docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "SecureMind-Architecture.docx"

NAVY = RGBColor(0x1B, 0x28, 0x38)
INDIGO = RGBColor(0x31, 0x2E, 0x81)
BODY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
RULE = "D1D5DB"
HEADER_FILL = "1B2838"
ALT_ROW = "F4F6F8"
WHITE = "FFFFFF"


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run_font(run, name: str, size: int, color: RGBColor, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_text(paragraph, text: str, *, size=11, color=BODY, bold=False, italic=False, name="Calibri") -> None:
    run = paragraph.add_run(text)
    set_run_font(run, name, size, color, bold)
    run.italic = italic


def body_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, text)
    p.style = doc.styles["Normal"]


def purpose(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, "Purpose. ", size=11, color=INDIGO, bold=True)
    add_text(p, text)


def bullet(doc: Document, lead: str, rest: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    add_text(p, "•  ", size=11, color=INDIGO, bold=True)
    add_text(p, lead, bold=True)
    add_text(p, rest)


def h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        set_run_font(run, "Calibri", 16, NAVY, True)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)


def h2(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        set_run_font(run, "Calibri", 13, INDIGO, True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def set_cell_border(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), RULE)
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def write_cell(cell, text: str, *, header: bool = False, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    color = RGBColor(0xFF, 0xFF, 0xFF) if header else BODY
    add_text(p, text, size=10, color=color, bold=header or bold)
    set_cell_border(cell)
    if header:
        shade(cell, HEADER_FILL)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = False
    for i, h in enumerate(headers):
        write_cell(table.rows[0].cells[i], h, header=True)
        table.rows[0].cells[i].width = Inches(widths[i])
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            write_cell(table.rows[r + 1].cells[c], value, bold=(c == 0))
            table.rows[r + 1].cells[c].width = Inches(widths[c])
            if r % 2 == 1:
                shade(table.rows[r + 1].cells[c], ALT_ROW)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    r = run._r

    def fld(kind: str) -> None:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r.append(el)

    fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    r.append(instr)
    fld("separate")
    text = OxmlElement("w:t")
    text.text = " "
    r.append(text)
    fld("end")


def header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(hp, "SECUREMIND", size=9, color=NAVY, bold=True)
    add_text(hp, "   ·   Architecture", size=9, color=MUTED)
    tab = hp.paragraph_format
    tab.space_after = Pt(2)
    # right-aligned draft mark via a second run is awkward; use a tab stop
    p_pr = hp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "right")
    tab_el.set(qn("w:pos"), "9360")
    tabs.append(tab_el)
    p_pr.append(tabs)
    hp.add_run("\t")
    add_text(hp, "Draft  0.1", size=9, color=MUTED)

    # bottom border on header paragraph
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1B2838")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_pr = fp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "right")
    tab_el.set(qn("w:pos"), "9360")
    tabs.append(tab_el)
    p_pr.append(tabs)
    top = OxmlElement("w:pBdr")
    top_el = OxmlElement("w:top")
    top_el.set(qn("w:val"), "single")
    top_el.set(qn("w:sz"), "6")
    top_el.set(qn("w:space"), "1")
    top_el.set(qn("w:color"), RULE)
    top.append(top_el)
    p_pr.append(top)
    add_text(fp, "Official architecture document", size=9, color=MUTED)
    fp.add_run("\t")
    add_text(fp, "Page ", size=9, color=MUTED)
    add_field(fp, " PAGE ")
    add_text(fp, " of ", size=9, color=MUTED)
    add_field(fp, " SECTIONPAGES ")


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = normal.paragraph_format
    pf.line_spacing = 1.12
    pf.space_after = Pt(8)


def build() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.9)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    section.different_first_page_header_footer = True

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    add_text(kicker, "SECUREMIND", size=12, color=INDIGO, bold=True)
    kicker.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    add_text(title, "Architecture", size=36, color=NAVY, bold=True, name="Calibri")

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    add_text(sub, "The governing description of the application.", size=14, color=MUTED)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "312E81")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(2)
    add_text(meta, "Document status", size=9, color=MUTED, bold=True)
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(10)
    add_text(line, "Draft 0.1  ·  26 September 2026  ·  For review", size=12, color=NAVY)

    note = doc.add_paragraph()
    add_text(
        note,
        "This document leads the rest of the documentation. Product behavior described here is the behavior to keep. A later PDF is produced from this file after review.",
        size=11,
        color=BODY,
    )

    # Body section so the cover has no header
    body = doc.add_section()
    body.page_width = Inches(8.5)
    body.page_height = Inches(11)
    body.left_margin = Inches(1)
    body.right_margin = Inches(1)
    body.top_margin = Inches(1.05)
    body.bottom_margin = Inches(0.9)
    body.header_distance = Inches(0.4)
    body.footer_distance = Inches(0.4)
    body.different_first_page_header_footer = False
    header_footer(body)

    # Restart page numbers at 1 for the body
    sect_pr = body._sectPr
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sect_pr.append(pg)

    h1(doc, "Document control")
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["Title", "SecureMind Architecture"],
            ["Status", "Draft 0.1 — for review"],
            ["Date", "26 September 2026"],
            ["Governs", "Product behavior, credentials, and the named screens"],
            ["Does not govern", "Marketing pages, or a public hosted service"],
            ["Companion drawings", "docs/diagrams/securemind-onprem.drawio and securemind-aws.drawio"],
        ],
        [1.7, 4.8],
    )
    body_para(
        doc,
        "Read this document first. When a screen and this document disagree, treat the disagreement as a defect to resolve, not as two truths. Comments from review are applied here, then the application is changed to match.",
    )

    contents = doc.add_paragraph()
    contents.paragraph_format.space_before = Pt(16)
    contents.paragraph_format.space_after = Pt(6)
    add_text(contents, "Contents", size=16, color=NAVY, bold=True)
    toc_note = doc.add_paragraph()
    toc_note.paragraph_format.space_after = Pt(4)
    add_text(toc_note, "Right-click the contents and choose Update field if the page numbers are blank.", size=10, color=MUTED, italic=True)
    toc = doc.add_paragraph()
    add_field(toc, ' TOC \\o "1-2" \\h \\z \\u ')

    h1(doc, "1.  How to read this document")
    body_para(
        doc,
        "Each part of the application is one section. The section names what is on the screen, what a person can do, and why the part exists. Where one idea covers several controls, they stay together. Fixed model is part of Models. It is not its own chapter.",
    )
    body_para(
        doc,
        "The audience is the person who will operate an install, and the person who will change the product. The language is the language of the screens.",
    )

    h1(doc, "2.  Legend")
    body_para(doc, "Marks and words used in the application and in this document.")
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Install", "One running SecureMind: Docker on a machine you control, or your own AWS account."],
            ["Administrator", "The person who connects model sources, sets detection policy, creates accounts, and sets cost caps."],
            ["Member", "A person signed in to that install. Members use the administrator’s model sources."],
            ["Workspace", "The boundary for chat, Knowledge, memory, and saved items. One workspace does not read another."],
            ["Complexity tiers", "Max, Plus, Standard, Light, Micro. Chat picks the tier. You assign the model on each row."],
            ["Fixed model", "Every message in the workspace uses the one model you pick."],
            ["Enable", "Check one cloud model. If it answers, it joins the chat list."],
            ["Check access", "Run that check across the models listed for one vendor key."],
            ["Test", "Recheck a model that is already in the chat list."],
            ["Hide / Show", "Remove a model from the chat list, or put it back. A hidden model can still sit on a tier."],
            ["Hidden", "Yellow mark. The assigned model was hidden from the chat list. The tier can still run it."],
            ["Check failed", "The model did not answer when it was checked. It stays out of chat."],
            ["Blue count", "How many replies that tier ran in the selected time range."],
            ["Provider chip", "How many models of that kind can be assigned. Example: Gemini 9. Not a call count."],
            ["Vendor key", "A Groq, Gemini, OpenAI, or Anthropic key saved encrypted on this install."],
            ["Bedrock access", "The AWS account already on the machine. An access key is never typed in and never stored."],
            ["Cost estimate", "Recorded tokens times list rates, plus 10 percent, rounded up. Not an invoice."],
            ["Detection", "A pattern scan before send and at document ingest. Risk reduction, not a guarantee."],
        ],
        [1.6, 4.9],
    )

    h1(doc, "3.  Purpose")
    purpose(
        doc,
        "SecureMind is an AI workspace a team runs for itself: model choice, a review before send, workspace memory, Knowledge with sources, and a visible estimate of use, cost, and what the scanner flagged.",
    )
    body_para(
        doc,
        "The application is self-contained. The administrator connects the model sources for that install. Everyone who signs in uses those sources. Members do not bring a private cloud key for SecureMind to store on a server operated by someone else.",
    )
    body_para(
        doc,
        "SecureMind is not offered as a live service that holds keys for the general public. The intended path is to run it yourself. Clone it and run it with Docker, or host it in an AWS account you control. The keys, the documents, and the chat history stay with that environment.",
    )
    body_para(
        doc,
        "Detection reduces risk. It does not guarantee that every sensitive value is found, and it is not a compliance program. The person sending a message still reviews it.",
    )

    h1(doc, "4.  Credentials")
    purpose(
        doc,
        "One install, one set of model sources, controlled by the administrator.",
    )
    bullet(doc, "Vendor keys. ", "Groq, Gemini, OpenAI, and Anthropic keys are saved on the install, encrypted. Only an administrator can write them. The screen shows a masked hint, never the secret.")
    bullet(doc, "Bedrock. ", "Billable AWS is turned on by the administrator. The credential is the AWS chain already present: an SSO login, a shared profile, environment variables, or a role. An AWS access key typed into Settings is refused and is not stored.")
    bullet(doc, "Local models. ", "Ollama runs on the same machine. It does not require a vendor key.")
    bullet(doc, "Shared use. ", "Members of the install chat with those sources. A monthly cost cap, set per person, limits cloud spend estimated for that person.")
    body_para(
        doc,
        "This separation is the reason the product is distributed as something you host. A public service that stored every customer’s key would be a different product, and it is not this one.",
    )

    h1(doc, "5.  Two ways to run it")
    purpose(doc, "The same application, two places for the data.")
    h2(doc, "5.1  On your machine")
    body_para(
        doc,
        "Docker Compose runs the application, Postgres with pgvector, a local sign-in stand-in, and Ollama. Chat, Knowledge, and memory stay in that Postgres database. Optional Bedrock uses a read-only mount of the computer’s AWS configuration. Optional vendor calls use the key the administrator saved. Uploads, profile photos, and the key-encryption secret are files on that machine.",
    )
    h2(doc, "5.2  In your AWS account")
    body_para(
        doc,
        "The interface is hosted on Amplify. Chat streaming goes through API Gateway to a Lambda. Sign-in is an invite-only Cognito user pool. Workspaces, messages, and vectors live in DynamoDB. Vendor keys live in Parameter Store as an encrypted SecureString. Profile photos live in a private S3 bucket. There is no custom VPC. These are regional services in the account you operate.",
    )
    body_para(
        doc,
        "A push to the application repository can test the code, apply Terraform, and start an Amplify build when that account’s deploy role is configured. The public marketing pages are a separate site and are outside this document.",
    )

    h1(doc, "6.  Workspaces")
    purpose(
        doc,
        "Keep one body of work together, and keep it out of the next one.",
    )
    body_para(
        doc,
        "A workspace owns its chat, Knowledge, memory, and saved items. A search in one workspace does not read another. The name you give the workspace is the name shown in chat. The emoji and color distinguish workspaces in the sidebar. They do not choose a model and they are not a personality.",
    )
    body_para(
        doc,
        "Switching workspace loads that conversation. A message queued and not yet sent is dropped, so it is not delivered into the wrong workspace.",
    )

    h1(doc, "7.  Chat")
    purpose(
        doc,
        "Write, review, and receive a reply, with the model choice and the workspace boundary already decided.",
    )
    body_para(
        doc,
        "The composer stays open while a reply is generating. The next message can be typed and queued. It waits, already reviewed, and sends when the current reply finishes. Stop ends the reply in progress. Attachments stay with the queued message.",
    )
    body_para(
        doc,
        "Before a message leaves the browser, a scan looks for sensitive patterns. If something matches, you see it and decide what continues. The text that is stored and sent is the reviewed text. The scan can miss a value. A request built outside the normal screen can skip that review, so the browser step is a control for the person at the keyboard, not a boundary against a crafted call.",
    )
    body_para(
        doc,
        "The answering model comes from Models. Complexity tiers let the message select Max, Plus, Standard, Light, or Micro, and each tier runs the model assigned to that row. Fixed model sends every message to the one model you picked. Tiers are the main path.",
    )
    body_para(
        doc,
        "A reply names Knowledge sources when the question is about material in the workspace and a source was actually used. A greeting or a general question does not invent a source. Facts kept for later come from what you wrote, not from the assistant’s replies.",
    )

    h1(doc, "8.  Knowledge")
    purpose(doc, "Let an answer point at documents that belong to this workspace.")
    body_para(
        doc,
        "Knowledge is files you upload, and files you attach in chat. Text is scanned as it is ingested, before it is stored or embedded. A later answer can cite a passage that was retrieved. A general question does not pretend a document said something it did not.",
    )

    h1(doc, "9.  Memory")
    purpose(doc, "Remember the statements you meant to keep, and let you take one back.")
    body_para(
        doc,
        "Memory holds facts you can correct, pin, or forget, and a short summary of the current conversation so a long thread can continue. Forgotten facts stay stored and are hidden from chat. The assistant’s replies are not saved as facts.",
    )

    h1(doc, "10.  Saved")
    purpose(doc, "Keep a line without turning it into a fact or a document.")
    body_para(
        doc,
        "Saved items are passages you mark from a conversation. An emoji in the text can show on the saved row so the line is recognizable in the list.",
    )

    h1(doc, "11.  Insights")
    purpose(
        doc,
        "Show, for this account, which models ran, what they cost, and what the scanner flagged.",
    )
    body_para(
        doc,
        "Insights sits in the rail beside chat. Each tab is a summary. Opening the full view leaves the rail for the same subject with charts, a time range, and a workspace filter. An administrator can open those full pages for one member from the Admin console.",
    )

    h2(doc, "11.1  Models")
    body_para(
        doc,
        "This panel is the complexity-tier assignment. Chat picks a tier from the message. Each row is the model that tier will run. Click a row to assign a different one. A cloud model appears as a choice after its source is connected and Enable has succeeded.",
    )
    body_para(
        doc,
        "The blue number on a row is how many replies that tier ran in the selected range. Today, Weekly, Monthly, and All time change that count, and so does the workspace filter. The provider chips, such as Gemini 9, count models you can assign. They are not call counts.",
    )
    body_para(
        doc,
        "A yellow Hidden mark means that model was hidden under Settings → Models. It is off the chat model list and can still be the model a tier runs. Show puts it back on the list. Check failed means the model did not answer when it was checked.",
    )
    body_para(
        doc,
        "Fixed model is the other position of the same switch. Every message uses the one model you pick. It is the ordinary single-model setup. The tiers are the path this application is built around.",
    )
    body_para(
        doc,
        "Open full usage view opens the Usage page. Open full-size editor in Settings opens the same assignment on the Models page.",
    )

    h2(doc, "11.2  Usage")
    body_para(
        doc,
        "Usage is the record of replies: how many ran, which tier, which model, how long they took, and whether an answer was grounded in Knowledge. Filter by time and by workspace. The rail shows the current assignment and a count. Usage is where you look across the install.",
    )

    h2(doc, "11.3  Cost")
    body_para(
        doc,
        "Cost estimates spend from the tokens each reply recorded. Groq and Bedrock use list rates. A local Ollama model is zero. The estimate is raised by 10 percent and rounded up, so it tends high. It is not an AWS invoice, and Cost Explorer is not read. An administrator can edit the rates used for new replies. Older replies keep the estimate they were given.",
    )
    body_para(
        doc,
        "The page compares that spend with what the same replies would have cost if every one had used the most expensive default tier. The gap illustrates routing. It is not money returned. A monthly cap, set per person in the Admin console, warns as it is approached and can block new cloud replies at the cap. The cap uses this estimate.",
    )

    h2(doc, "11.4  Security")
    body_para(
        doc,
        "Security is the log of what the scanner flagged: the category, whether it was in chat or in a document, and what was done with it. The matched text is not in the log. You can dismiss or delete an entry. Deleting the entry does not restore the original value.",
    )
    body_para(
        doc,
        "Sensitivity chooses which pattern families run from now on. Low runs the highest-signal families, including email, national identifiers, card numbers, and secrets. Balanced adds phone numbers, IP addresses, and street addresses. High adds passport-like patterns. A change does not rescan old messages or old documents.",
    )
    body_para(
        doc,
        "Safe terms are exemptions for values you know are not personal, such as an internal IP range or a product name. Social Security numbers, Social Insurance numbers, card numbers, and secrets cannot be exempted. Safe terms are edited with the instance detection policy.",
    )

    h1(doc, "12.  Models settings")
    purpose(
        doc,
        "Decide which accounts and which local models this install may call, prove a cloud model can answer, and assign a model to each tier.",
    )
    body_para(doc, "Settings → Models has three layers.")
    bullet(doc, "Connections. ", "Vendor keys, Amazon Bedrock, Ollama, and voice. Voice stays off until Groq Whisper or Amazon Transcribe is chosen, with an acknowledgement that the audio is sent before any text scan. The transcript lands in the composer and is reviewed on send.")
    bullet(doc, "Model list. ", "Everything those sources offered. A cloud model stays under Not in chat yet until Enable succeeds. Check access tries the list for one key. Hide and Show control the chat list.")
    bullet(doc, "Complexity tiers and Fixed model. ", "The same control as the Insights rail, with more room. Enabling Bedrock does not assign a model and does not change rows already set.")

    h1(doc, "13.  Detection policy")
    purpose(doc, "One scan rule for the install, set by the person who operates it.")
    body_para(
        doc,
        "Settings → Workspace holds sensitivity and safe terms for the whole install. It is administrator-owned. It is separate from a workspace’s name and description, which are edited on that workspace. A member does not supply a private exemption list that the server must obey.",
    )

    h1(doc, "14.  My Profile")
    purpose(doc, "Identify the signed-in person, and let them change their own password.")
    body_para(
        doc,
        "Profile holds a display name, a preset picture or a cropped photo, and a password change when the install uses its directory. The photo is stored for this install and served only to a signed-in session. It is not a public link.",
    )

    h1(doc, "15.  Appearance")
    purpose(doc, "Make the same install readable in a dark room or a bright one.")
    body_para(
        doc,
        "Dark or light, for the whole application. Dark is the default. The choice is saved on this device and applies immediately. It does not change anyone else’s screen.",
    )

    h1(doc, "16.  Admin console")
    purpose(
        doc,
        "Add the people who share this install, limit their spend, and look at one person without mixing them into everyone else.",
    )
    body_para(
        doc,
        "An administrator creates an account, chooses a role, and sets a password. The password is not emailed. If mail is configured, an invite can point at the sign-in page. An account can be suspended, and a password can be reset. Each person has a monthly cost cap. A default cap applies when one is not typed.",
    )
    body_para(
        doc,
        "The console shows spend against those caps, and it opens Usage, Cost, and Security for the whole install or for one person.",
    )

    h1(doc, "17.  Getting started")
    purpose(doc, "Reach a first model. The rest of the application assumes that happened.")
    body_para(
        doc,
        "Getting started is a short pictured sequence for connecting a model. It is linked from the sidebar while chat has no model ready. It is not a tour of every screen in this document.",
    )

    h1(doc, "18.  How the parts fit")
    body_para(
        doc,
        "You run the install. The administrator connects keys and Bedrock, sets the detection rule, creates people, and sets caps. Each person works inside a workspace. Chat reviews a message, then a tier or a fixed model answers. Knowledge and memory stay in that workspace. Insights shows what ran, what it cost, and what the scanner flagged.",
    )
    body_para(
        doc,
        "Nothing in that chain requires a SecureMind-operated service to hold your keys.",
    )

    h1(doc, "19.  What this document does not claim")
    bullet(doc, "Complete removal. ", "The scan can miss a value.")
    bullet(doc, "A compliance program. ", "Categories and a log are not a certification.")
    bullet(doc, "An invoice. ", "Cost figures are estimates from list rates.")
    bullet(doc, "A public multi-tenant service. ", "Each install is owned and keyed by the people who run it.")
    bullet(doc, "Stored AWS access keys. ", "Bedrock does not accept a key pasted into the application.")

    # Ask Word to refresh TOC/page fields on open
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
