"""Build the SecureMind technical guide Word document.

Regenerate after edits to this script:

    python scripts/build-technical-guide-docx.py

Regeneration replaces the Word file. Edits made only inside Word are lost.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "SecureMind-Technical-Guide.docx"

INK = RGBColor(0x1A, 0x1C, 0x1F)
MUTED = RGBColor(0x5C, 0x61, 0x6B)
VIOLET = RGBColor(0x5B, 0x4D, 0xB8)
VIOLET_INK = RGBColor(0x3A, 0x32, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

PRIMARY = "5B4DB8"
CARD = "F4F5F7"
RULE = "D0D3DA"


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run_font(run, name: str, size: int, color: RGBColor, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_text(paragraph, text: str, *, size=11, color=INK, bold=False, italic=False, name="Calibri") -> None:
    run = paragraph.add_run(text)
    set_run_font(run, name, size, color, bold)
    run.italic = italic


def body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text)


def h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        set_run_font(run, "Calibri", 18, VIOLET, True)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)


def h2(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        set_run_font(run, "Calibri", 13, VIOLET_INK, True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def set_cell_border(cell, color: str = RULE) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def set_cell_margins(cell, twips: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def write_cell(cell, text: str, *, header: bool = False, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    color = WHITE if header else INK
    add_text(p, text, size=10, color=color, bold=header or bold)
    set_cell_border(cell)
    set_cell_margins(cell)
    if header:
        shade(cell, PRIMARY)
    elif fill:
        shade(cell, fill)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = False
    table.allow_autofit = False
    for i, heading in enumerate(headers):
        write_cell(table.rows[0].cells[i], heading, header=True)
        table.rows[0].cells[i].width = Inches(widths[i])
    for r, row in enumerate(rows):
        fill = CARD if r % 2 == 0 else "EEF0F3"
        for c, value in enumerate(row):
            write_cell(table.rows[r + 1].cells[c], value, bold=(c == 0), fill=fill)
            table.rows[r + 1].cells[c].width = Inches(widths[c])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    r = run._r

    def fld(kind: str) -> None:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r.append(el)

    fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    r.append(instr)
    fld("separate")
    text = OxmlElement("w:t")
    text.text = " "
    r.append(text)
    fld("end")


def header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    p_pr = hp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "right")
    tab_el.set(qn("w:pos"), "9360")
    tabs.append(tab_el)
    p_pr.append(tabs)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), PRIMARY)
    border.append(bottom)
    p_pr.append(border)
    hp.paragraph_format.space_after = Pt(4)
    add_text(hp, "SECUREMIND", size=9, color=VIOLET, bold=True)
    add_text(hp, "   ·   Technical guide", size=9, color=MUTED)
    hp.add_run("\t")
    add_text(hp, "Version  1.2", size=9, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    p_pr = fp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "right")
    tab_el.set(qn("w:pos"), "9360")
    tabs.append(tab_el)
    p_pr.append(tabs)
    top = OxmlElement("w:pBdr")
    top_el = OxmlElement("w:top")
    top_el.set(qn("w:val"), "single")
    top_el.set(qn("w:sz"), "6")
    top_el.set(qn("w:space"), "1")
    top_el.set(qn("w:color"), RULE)
    top.append(top_el)
    p_pr.append(top)
    add_text(fp, "Version 1.2", size=9, color=MUTED)
    fp.add_run("\t")
    add_text(fp, "Page ", size=9, color=MUTED)
    add_field(fp, " PAGE ")
    add_text(fp, " of ", size=9, color=MUTED)
    add_field(fp, " NUMPAGES ")


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.95)
    section.bottom_margin = Inches(0.85)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    banner = doc.add_table(rows=1, cols=1)
    cell = banner.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "SECUREMIND", size=11, color=WHITE, bold=True)
    t = cell.add_paragraph()
    t.paragraph_format.space_before = Pt(8)
    t.paragraph_format.space_after = Pt(2)
    add_text(t, "Technical Guide", size=28, color=WHITE, bold=True)
    s = cell.add_paragraph()
    s.paragraph_format.space_before = Pt(2)
    s.paragraph_format.space_after = Pt(6)
    add_text(s, "How the application decides, stores, and prices a turn", size=12, color=RGBColor(0xE4, 0xE2, 0xF2))
    shade(cell, PRIMARY)
    set_cell_border(cell, PRIMARY)
    set_cell_margins(cell, 160)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(10)
    meta.paragraph_format.space_after = Pt(8)
    add_text(meta, "Version 1.2", size=11, color=VIOLET, bold=True)
    add_text(meta, "   ·   25 October 2026", size=11, color=MUTED)

    body(doc, "This guide explains how SecureMind works. The feature guide describes what a person uses. This one describes the decisions the software makes, the data that moves, and the rules behind complexity tiers and estimated cost.")
    body(doc, "It is written for developers, technical hiring managers, and the person operating an installation. File and function names are included so a reader can find the rule in the source. The prose stands on its own.")
    body(doc, "Detection reduces exposure. It does not certify that content is free of sensitive data, and it is not a compliance program. Cost figures are estimates from recorded tokens and a curated rate table. They are not invoices.")
    body(doc, "Version 1.1 reflected the October 2026 security assessment and its fixes: authentication that fails closed, a server-side check for the highest-risk values, sign-in and upload limits, an audit log, a locked-down Docker stack, and monitoring and recovery on AWS. The assessment itself is in docs/security-audit/.")
    body(doc, "Version 1.2 adds the changes made since: live voice captions straight from the browser to Amazon Transcribe, voice counted toward the cost cap, files built from chat, an approved helper model for them, automatic unlisting of retired models, a score card that explains every tier decision, evidence-aware tier floors, a wider server-side check, a real cap on how far an uploaded document may unpack, and a read-only role for pull-request plans.")

    h1(doc, "1  How to read this guide")
    body(doc, "Read section 3 first. It is the path of one chat message. Later sections expand one step of that path. Complexity routing is section 8. Cost is section 15.")
    body(doc, "Two installations exist. Local means Docker on the operator's machine, or a developer process, with Postgres. Hosted means the operator's own AWS account, with DynamoDB. Behavior that differs is called out in the section where it matters, and again in section 18.")

    h1(doc, "2  What the system is")
    body(doc, "SecureMind is a single-operator AI workspace. One installation has its own users, vendor keys, detection policy, and stored conversations. It is not a shared public service that holds every customer's keys.")
    body(doc, "One installation serves one organisation. It is recommended for 1 to 25 people and can serve up to about 100 with the adjustments in section 19. Users' data is kept apart from each other, but Admins act across the whole installation, so separate organisations should each run their own installation.")
    body(doc, "The useful layer is detection before a typed message is sent, workspace memory, Knowledge with citations, complexity-tier model choice, and a visible estimate of spend and detected-risk. The product name in the chrome is SecureMind. The name shown as the speaker in a chat is the workspace name. That name is a label. It is not a personality, and it is not injected into social or general turns.")
    body(doc, "A browser-supplied model id is a preference. The server accepts it only through resolveAllowedModel() in src/lib/models/allowlist.ts. Anything else is ignored and the router's choice stands. The server does not infer a provider from a substring of a client-supplied id.")

    h1(doc, "3  One chat turn")
    body(doc, "The chat path is src/lib/chat/run-chat-stream.ts. Locally the browser calls the Next.js route on the same origin. On the hosted stack the stream goes to API Gateway and the chat Lambda. The UI and the other APIs stay on Amplify. Chat is not streamed through an Amplify Next route.")
    add_table(
        doc,
        ["Step", "What happens"],
        [
            ["1. Composer", "The user can type the next message while a reply is still generating. The queue holds the already-reviewed text and a snapshot of its attachments. Stop aborts the current stream. Switching workspace clears the queue and the composer."],
            ["2. Detection", "For typed text, pattern detection runs in the browser. A match opens a review dialog. The text that continues is the reviewed replacement. There is no UI path that sends the original matched value. On arrival, the server checks every text field again for the categories that can never be exempted (step 4b)."],
            ["3. Identity", "The server takes the user id from the session (getIdentity). It does not read a user id from the body, the query string, or a client header. A missing, forged, expired or wrong-client token is refused, and a suspended or deleted account is refused even while its token is still valid."],
            ["4. Cap", "If this user's UTC-month estimate is already at the cap, and the call is billable, the server returns 402 and does not call the model. Local Ollama is not in that cap. The voice routes make the same check before any audio is processed."],
            ["4b. Server check", "scrubChatRequestServerSide (src/lib/chat/server-scrub.ts) replaces any US Social Security number, Canadian SIN, card number or secret it finds in the message, history, attachments, summary or facts, so a request that skipped the browser cannot send them. The same gate, scrubServerText, covers text that is saved or sent to a model outside a chat turn (see section 6). Other categories stay a browser-side review decision."],
            ["5. Turn policy", "decideTurn classifies the message. That decision controls retrieval, citations, and whether workspace context is added to the prompt. It does not choose the model, except when it forces the Micro tier."],
            ["6. Tier", "routeMessage scores the message and picks Max, Plus, Standard, Light, or Micro, and writes a score card that explains the choice. A narrow ambiguous band may ask a small Groq model to break the tie. After retrieval, a turn answered from documents is raised to at least Standard or Plus."],
            ["7. Model", "The tier selects the assigned model. A fixed-model choice assigns that same model to every tier. An image forces the vision model. The allowlist can reject the assignment."],
            ["8. Prompt", "Conversational rules apply on every turn. Operational grounding and workspace context apply only when the turn needs them. Retrieved passages are fenced."],
            ["9. Stream", "The reply streams back. The stored user message is the sanitized text. After the answer, citation markers are checked against the sources that were actually packed."],
            ["10. Record", "An inference event stores the model, the tier, token counts, grounding, and the cost estimate. A detection finding stores the category and the outcome. It does not store the matched string."],
        ],
        [1.3, 5.4],
    )

    h1(doc, "4  Identity and access")
    body(doc, "userId comes from getIdentity(req) in src/lib/identity.ts and from nowhere else. Every read of a workspace, session, message, document, or memory fragment is scoped by that id. A record that exists but belongs to someone else is returned as 404, not 403, so the response does not confirm that the id exists.")
    body(doc, "Default local development can run with AUTH_MODE=local. A production build refuses that fixed local identity unless ALLOW_INSECURE_LOCAL_IDENTITY=true is set, and an unset fallback setting no longer turns a bad token into the local owner: authentication fails closed. The Docker installation uses a local stand-in for Cognito (Floci) with AUTH_MODE=cognito. A hosted installation uses an invite-only Cognito user pool with email usernames and the groups Owner, Admin, and Member. Roles are enforced on admin routes through requireAdmin. Members use the shared provider connections and cannot read or change the secret keys.")
    body(doc, "Tokens are checked for signature, issuer, token use (access tokens only) and the app client (src/lib/auth-claims.ts, shared by the edge middleware and getIdentity). On the hosted install, API Gateway also checks the Cognito token before the chat Lambda runs. Signing out ends the Cognito session, and disabling an account revokes its refresh tokens. getIdentity refuses a suspended or deleted account within about 30 seconds. An access token that was already issued stays valid until it expires, up to an hour, because it is checked by its signature rather than looked up on each request. Sign-in, the two-step code, password reset and the first-run claim are rate limited per address and per account (src/lib/rate-limit.ts); Cognito's own lockout is the backstop on AWS.")
    body(doc, "The first account on an empty pool is claimed through the setup flow with a one-time setup code: Terraform creates it on the hosted install, and the Docker init container prints it (docker compose logs init). The claim is reserved with a conditional write, so two simultaneous claims cannot both create an Owner. After that, an administrator creates accounts, assigns roles, sets passwords, and can suspend an account. Only the Super Admin can change another Admin's role, status, password or two-step verification. Every administrative change is written to an append-only audit log, read through /api/admin/audit. Delete removes sign-in. It does not wipe Knowledge or history, and a deleted account can be restored.")
    body(doc, "Two-step verification is optional authenticator-app (TOTP) sign-in through real Cognito, so it is a hosted-edition feature. Cognito makes the secret and checks every code. src/lib/totp.ts exists for tests and teaching and is never used to decide who may sign in. The routes under src/app/api/auth/mfa act only on the caller's own access token and take no user id from the request. Turning it off needs the password again and is rate limited per account. The Docker sign-in does not support codes, so the control is hidden there. Admin reset goes through adminResetDirectoryUserMfa, which respects the Super Admin protection.")

    h1(doc, "5  Workspaces, sessions, and the queue")
    body(doc, "A workspace is the boundary for chat history, Knowledge, facts, and bookmarks. Retrieval does not cross workspaces. One signed-in user can own many workspaces. The emoji and color are visual identity only. They are not sent to the model as a personality, and they do not select a model.")
    body(doc, "Each workspace has one continuous session, messages, an optional rolling summary, user-authored facts, documents, and bookmarks. It can also store a short description and an operational context string. That context is added to the prompt only when needsOperationalPrompt is true. Social and general turns do not receive the workspace name or mission.")
    body(doc, "An optional fixed model on the workspace, or the fixed-model control in Insights, bypasses per-tier model choice. The application still computes a tier, so Insights can count the reply against that tier. Every tier is then pointed at the same model id, provided the allowlist accepts it.")
    body(doc, "Switching workspace aborts an in-flight stream, drops the outbound queue, clears the composer, and loads that workspace's session. A queued message cannot be delivered into the wrong workspace.")

    h1(doc, "6  Detection")
    body(doc, "Detection is a pattern scan in scrubPII (src/lib/pii.ts), not a model. It redacts a span it recognizes. A detail outside those shapes is left as written. That is best effort. It is not a certification that the text is clear, and it is not a compliance program.")
    body(doc, "Chat runs it in the browser before the draft leaves the page. Knowledge ingest runs the same function on the server after text extraction and before the text is stored or embedded (scrubDocumentText in src/lib/documents.ts). The server reads sensitivity and safe terms from the administrator-owned instance policy. It does not take them from the upload. A picture is turned into text first, then scanned the same way. A voice transcript is scanned when it is sent. The model's reply is not scanned. A crafted request can skip the browser, so the client is not a security boundary. For that reason the server checks chat requests again, for the non-exemptible categories only (SSN, SIN, card number, secrets), and replaces any match before the model call. The same gate, scrubServerText, is applied to text that is saved or sent to a model outside a chat turn: a saved message (which also accepts only the roles user and assistant, and a bounded length), the transcript sent to summarize, workspace context, and a corrected memory fact. Emails, phone numbers, addresses and the other exemptible categories remain a browser-side decision, because the person may have reviewed and kept them on purpose.")
    body(doc, "When the browser finds a match, the user sees a review dialog. The raw matched value is not written to the database, localStorage, logs, or the model call. There is no restore path. Dismissing or deleting a finding removes the log entry. It does not put the original text back.")
    h2(doc, "How a match is decided")
    body(doc, "The function walks a fixed order. First it resolves the tier. Balanced is the default. A slider value at or below 0.25 is Low, at or below 0.60 is Balanced, and above that is High. A stored value of extreme is treated as High. Second, it keeps only the families whose minimum tier is at or below the active tier. Tiers are cumulative. Third, each family searches the whole text. An empty match is skipped. Fourth, a family with an extra check drops a hit that fails it. Fifth, if that family allows an exemption, safe terms are tested and a covered hit is dropped. The exemption is counted by term and category. The value is not recorded. Families that cannot be exempted skip this step. Sixth, surviving hits are sorted by where they start, and overlaps are merged into one span. The category kept is the one that started first. Seventh, each span is replaced with a label such as [email redacted]. The label names the category and contains none of the original characters. If nothing survived, the original text continues.")
    h2(doc, "Sensitivity")
    body(doc, "The administrator owns the policy in instance configuration. Balanced is the default. A change applies to new chat and new Knowledge only. Stored messages and documents are not re-scanned, because the original matched text was never kept.")
    add_table(
        doc,
        ["Family", "Runs from", "Can a safe term exempt it?"],
        [
            ["Email addresses", "Low", "Yes"],
            ["US Social Security numbers", "Low", "No"],
            ["Canadian social insurance numbers", "Low", "No"],
            ["Card numbers (Luhn check)", "Low", "No"],
            ["API keys and private keys", "Low", "No"],
            ["Phone numbers", "Balanced", "Yes"],
            ["IPv4 addresses", "Balanced", "Yes"],
            ["Street addresses", "Balanced", "Yes"],
            ["Passport and national-ID shapes", "High", "Yes"],
        ],
        [3.2, 1.5, 2.0],
    )
    body(doc, "What each extra check actually does: a US Social Security number is a 3-2-4 grouping with a hyphen or a space, and it has no checksum. A bare nine-digit run is not treated as one. A Canadian social insurance number is a 3-3-3 grouping or a bare nine-digit run, then a check digit taken from the left. All zeros and all-identical digits are rejected. A card number is 13 to 19 digits, with spaces or hyphens allowed between them, then the Luhn checksum from the right. Without Luhn, request IDs in a log would match. A secret matches a known form only: a private-key block, Stripe-style sk_live_ or sk_test_ (and the rk_ forms), GitHub ghp_, gho_, ghu_, ghs_, ghr_, or github_pat_, a Slack xox token, an AWS access-key prefix AKIA or ASIA, a named AWS secret or session-token assignment, an OpenAI-style sk- key, or a Google AIza key. A long random string is not enough. Phone numbers are a North American shape, including a bare ten-digit number that does not start with 0 or 1. An IPv4 address is four numbers from 0 to 255. A street address is a street number, then words, ending in a street word such as Street, Avenue, Road, or Drive. A passport-like value is one or two capital letters followed by seven to nine digits. That is a shape, not a lookup against a government record. The pattern text itself is not reproduced here.")
    h2(doc, "Safe terms")
    body(doc, "A safe term is tested after a shape matches, and only for a family that allows it. Literal means the whole match equals the term, ignoring case. CIDR means the IPv4 address falls inside the block. Domain means the match is that domain, or it ends on a real boundary such as .acme.com or @acme.com. notacme.com does not match acme.com. Prefix means the match starts with the term, after a trailing asterisk is removed. The non-exemptible categories above cannot be listed. A proposed term that itself looks like one of those categories is refused, so the list cannot store one.")
    body(doc, "The install ships with visible defaults for documentation mail (example.com, example.org, example.net, and the reserved endings test, invalid, localhost, and example) and for non-routable IPv4 (the private ranges 10.0.0.0/8, 172.16.0.0/12, and 192.168.0.0/16, loopback, link-local, and the three documentation ranges). An administrator can remove them. The server applies the instance policy on ingest. It does not accept a safe-term list from the browser. Members may propose a change. Only an administrator applies it.")
    body(doc, "What Insights stores is the category, the source (chat or document), the outcome, a count, and the sensitivity that was in force. The matched string is not stored.")

    h1(doc, "7  Turn policy")
    body(doc, "decideTurn in src/lib/turn-policy/decide-turn.ts runs before retrieval and before the prompt is built. It answers a different question from the complexity tier. The tier asks which model is warranted. Turn policy asks whether this message should search Knowledge, show Sources, or receive the workspace's operational context.")
    body(doc, "The classifier is deterministic. It normalizes the text, strips a short conversational prefix, then walks a fixed order: attachments, social and corrections, application help, grounded follow-ups, workspace or investigation cues, general questions, and a short unclear remainder. reasonCode is telemetry. It is never written into the assistant's reply.")
    add_table(
        doc,
        ["Kind of turn", "Knowledge search", "Sources", "Prompt"],
        [
            ["Social, joke, correction, or a request to drop the topic", "No", "No", "Conversational only. Micro tier."],
            ["How to use the application", "No", "No", "Conversational only. Micro tier."],
            ["Where did you learn that?", "No", "No", "Conversational only. Micro tier."],
            ["A terse fragment, such as a word or two with no cue", "Silent search first", "Only if a hit reaches the fused score 0.55", "A strong hit upgrades the turn and answers from the passages. A miss asks for clarification. Not forced to Micro."],
            ["Attachment, no compare request", "No. The attachment is the source.", "No workspace Sources", "Operational framing. The tier score still runs unless the instruction is almost empty, which forces Micro."],
            ["Attachment compared with workspace Knowledge", "Yes", "Yes, if cited", "Operational framing."],
            ["Names this workspace, a file, or an investigation", "Yes", "Only sources the answer cites", "Operational framing, including workspace context."],
            ["Follow-up after a cited or grounded turn", "Yes, scoped to prior sources", "Yes, if cited", "Operational framing."],
            ["General question", "Silent search", "Only if a hit reaches the fused score 0.55", "Conversational unless a hit upgrades the turn."],
        ],
        [2.2, 1.6, 1.4, 1.5],
    )
    body(doc, "A general question does not show Sources by default. If the silent search returns a hit whose fused score reaches SPECULATIVE_MIN_SCORE (0.55), the turn is upgraded to workspace-grounded and Sources become available. A miss stays general, with no Sources and no announcement that Knowledge was empty. Workspace-grounded and investigation turns accept hits at GROUNDED_MIN_SCORE (0.40).")
    body(doc, "needsOperationalPrompt is true for retrieval, investigation, workspace-grounded turns, grounded follow-ups, and attachments. Those turns receive operational instructions and the workspace context. Other turns receive only the conversational base, so a greeting is not wrapped in the workspace mission.")
    body(doc, "When preferMicroTier is set, the chat path uses Micro even if the complexity score would have chosen a higher tier. That flag is set for social turns, application help, provenance questions, corrections, jokes, requests to leave a topic, short deictic phrases with nothing to look up (such as is this correct), and an attachment that arrives with almost no instruction. A terse fragment that may be a document lookup is deliberately not forced to Micro, because an upgraded turn must be answered by a model that can read the passages. Images are a separate override and are described in section 9.")

    h1(doc, "8  Complexity tiers")
    body(doc, "A tier chooses which model handles the turn. It does not change the assistant's personality. The five tiers, from most capable to least, have these default Groq assignments. An administrator can assign a different model to a tier. The assignment is kept only when that model is allowed to run.")
    add_table(
        doc,
        ["Tier", "Default model", "List rate, USD per million tokens"],
        [
            ["Max", "llama-3.3-70b-versatile", "0.59 in / 0.79 out"],
            ["Plus", "qwen/qwen3-32b", "0.29 in / 0.59 out"],
            ["Standard", "openai/gpt-oss-120b", "0.15 in / 0.60 out"],
            ["Light", "meta-llama/llama-4-scout-17b-16e-instruct", "0.11 in / 0.34 out"],
            ["Micro", "llama-3.1-8b-instant", "0.05 in / 0.08 out"],
        ],
        [1.3, 3.4, 2.0],
    )
    body(doc, "Standard is the default tier. Older stored events use High for Max and Fast for Light (normalizeStoredTier). The rates above are the curated Groq table in src/lib/pricing/groq.ts, stamp groq-list-2026-09, effective 22 September 2026. They are not read from a live pricing API. Section 15 explains how a turn's tokens become a dollar estimate.")
    h2(doc, "When a tier is chosen")
    body(doc, "Routing lives in src/lib/models/route-message.ts. It runs on the server for each send. The inputs are the current message, up to four recent messages, the tier of the previous assistant turn, and whether the turn includes an image. The decision is local first. A network call happens only for one narrow case.")
    body(doc, "Pass 1, trivial acknowledgements. A message that is only thanks, ok, yes, no, a short laugh, bye, or a short greeting of at most 48 characters is Micro immediately. The scorer and the classifier are skipped. An image is never treated as trivial.")
    body(doc, "Pass 2, a score from 0 to 100. The baseline is 42, which maps to Standard. Signals add or subtract, and the score is clamped to that range. Every number comes from a named signal in one table (SIGNAL_RULES), so the rule is the table and nothing else.")
    add_table(
        doc,
        ["Signal", "Effect"],
        [
            ["20 characters or fewer", "Subtract 18"],
            ["21 to 50 characters", "Subtract 8"],
            ["51 to 120 characters", "Subtract 2"],
            ["281 to 450 characters", "Add 8"],
            ["Longer than 450 characters", "Add 14"],
            ["Deep-analysis phrasing (a full design, a formal proof, a comprehensive architecture)", "Add 45"],
            ["Technical work longer than 40 characters (debug, refactor, stack trace, root cause, and similar)", "Add 28"],
            ["Code present, and the message is longer than 30 characters", "Add 32"],
            ["Planning language (help me plan, pros and cons, roadmap, weigh options)", "Add 20"],
            ["Comparing options or trade-offs", "Add 16"],
            ["Asking for decision advice, in a message longer than 55 characters", "Add 14"],
            ["Several questions or list items in a message longer than 60 characters", "Add 8"],
            ["Creative writing longer than 50 characters", "Add 10"],
            ["Casual chat, such as how are you or what do you think", "Add 8, which keeps ordinary conversation on Standard"],
            ["An emotional word", "Add 4"],
            ["A file attached with no instruction", "Add 8"],
            ["A how-to or an explain, 45 to 319 characters", "No points, but the score is capped at 28, which is Light"],
            ["A quick fact under 120 characters (who, when, define, capital of)", "Subtract 22, then cap the score at 18, which is Light"],
            ["An acknowledgement that is not a bare one-liner", "Subtract 35"],
            ["A greeting in an ongoing conversation (hello again, good morning)", "Raise a low score to 40, so it stays on Standard"],
        ],
        [3.2, 3.5],
    )
    body(doc, "The score maps to a tier at fixed cutoffs: 66 and above is Max, 52 to 65 is Plus, 36 to 51 is Standard, 12 to 35 is Light, and below 12 is Micro.")
    body(doc, "Pass 3, carry-over. A short message, under 160 characters, can inherit the previous tier minus one level, so that \"and the timeout?\" after a deep reply does not drop to a cheap model, and a thread decays back to Standard instead of staying expensive. Only Plus and Max carry, the carried tier never goes below Standard, and it does not apply when the scorer already marked the message as casual chat, emotional, a greeting again, a quick fact, a how-to, or an acknowledgement. A short follow-up after a Max reply therefore runs on Plus, and after a Plus reply on Standard.")
    body(doc, "Pass 4, a quality floor. After carry-over, a non-trivial message is not left on Micro. A quick fact or a how-to is raised to Light. Anything else that still scored Micro is raised to Standard. The only messages that stay on Micro through this function are the trivial acknowledgements from pass 1. Turn policy can still force Micro afterward, as described in section 7. That force happens in the chat path after routeMessage returns, so a social turn does not get lifted back to Standard by this floor.")
    body(doc, "Pass 5, evidence floors. After retrieval, applyEvidenceFloors raises the tier, never lowers it. A turn answered from the user's documents or from an attachment is at least Standard, because reading evidence needs a careful reader even when the words are short. A summary or analysis of documents (summarize, compare, risks, review, audit, differences, and similar), and any answer that draws on three or more documents, is at least Plus. The model is then resolved again for the new tier. If that tier has no usable model, or the call would breach the cost cap, the first-pass model stands and the score card says so.")
    body(doc, "The score card (src/lib/models/route-scorecard.ts) records the baseline, each signal with its points, every cap or adjustment, whether the tie-break ran and what it cost, and the final tier. It holds no message text. The reply shows it as a Why this tier? link. The plain-language meaning of each tier lives only in src/lib/models/tier-definitions.ts and feeds both the interface and the tie-break prompt, and a test checks that it matches the tier list in router.ts. A routing fixture in src/lib/models/route-fixtures.ts is added whenever a message lands on the wrong tier.")
    h2(doc, "The classifier")
    body(doc, "If the score is between 34 and 46 inclusive, and none of the clear-intent signals fired, the score is treated as ambiguous. Clear intent includes casual chat, emotion, greetings, facts, how-to, planning, comparison, decision advice, technical language, code, deep analysis, creative writing, several questions in one message, and a file attached with no instruction. A plain mid-length sentence with only the baseline, and no stronger signal, is the usual case that qualifies.")
    body(doc, "The tie-break is classifyComplexityLlm in src/lib/models/groq.ts. It calls llama-3.1-8b-instant at temperature 0, with at most ten tokens of output, and asks for one tier name. Up to four recent lines are included, each truncated. If Groq is not configured, or the call fails, the synchronous tier stands. If the classifier answers Micro for a non-trivial message, the result is raised to Standard. Carry-over is applied again to the classifier's choice. The call's token cost is recorded on the turn as extra cost, so it counts toward the monthly cap without creating a separate event that would inflate request counts. High and Fast, if the small model emits the old names, are mapped to Max and Light.")
    h2(doc, "What the tier does not do")
    body(doc, "The tier does not decide whether to search Knowledge. It does not add the workspace mission. It does not change detection. Background summarize and extract always use the cheapest configured model, llama-3.1-8b-instant when Groq is available, regardless of the tier of the conversation they are summarizing.")

    h1(doc, "9  The model that actually runs")
    body(doc, "After the tier is known, routeModel(\"chat\", tier) in src/lib/models/router.ts returns the built-in Groq model for that tier. The chat path then looks at the assignment map sent with the request. In fixed-model mode that map sets every tier to the one selected id. In tier mode it sets each tier to the model shown on that Insights row. resolveAllowedModel accepts the id only when it is a built-in router model, a known Bedrock text fallback, or a catalog entry that is usable for chat and has not failed verification. A rejected id is logged and ignored. The built-in model for the tier remains.")
    body(doc, "An image skips that assignment. The call is forced to meta-llama/llama-4-scout-17b-16e-instruct on Groq. The same model is the default for turning an uploaded image into text during Knowledge ingest, unless BEDROCK_VISION_MODEL or GROQ_VISION_MODEL is set. Image extraction on Bedrock defaults to amazon.nova-lite-v1:0.")
    body(doc, "If the chosen provider cannot run, the runtime decision can fall back to an installed Ollama chat model on a local installation. The application does not pull model weights on a chat request. Docker's first boot pulls nomic-embed-text and llama3.2:3b into the Ollama volume. Those local models use the machine's own resources and are slow on most computers, so a vendor API key (or Bedrock on AWS) is strongly recommended. A hosted installation has no Ollama. If no model is assigned and neither Groq nor Bedrock can run, hosted chat returns 503 and points the operator at Settings.")
    h2(doc, "Enable, Test, and Check access")
    body(doc, "A model discovered from a vendor key is not in the chat menu until it has answered a short check. That first success is Enable. It stores a verification, on the order of seven days, and the model moves into the chat list. Models already in the list show Test, which is the same check run again. Check access runs the check across the models listed for one provider key. Hide removes a model from the chat menu. A hidden model can still be the model a tier runs. Check failed means the last check did not succeed. It does not, by itself, explain the vendor's reason.")

    body(doc, "A cloud model that worked when it was verified can be retired by its provider before that verification expires. When a chat call fails, retireModelAfterChatFailure (src/lib/models/retire-failed-model.ts) reads the provider's error with the same classifier the Test button uses. If the answer is permanent (the model is gone, is not a chat model, or is not available to this key), it records a failed verification for that key, refreshes the model list, and adds a plain note to the reply saying the model was taken off the list. Temporary problems, such as a rate limit, a billing limit, a timeout or an outage, change nothing, and a missing local Ollama model is fixed by pulling it, so those are never unlisted. The function never throws, so a failure to record it cannot break the error path.")

    h1(doc, "10  Providers and credentials")
    body(doc, "Groq is the complexity-tier default when an administrator has stored a key, or when a local environment already has one. It is also used for background tasks, optional vision extraction, the tier classifier, and optional Whisper. Ollama is the local chat fallback and the local embedding model. Amazon Bedrock is optional.")
    body(doc, "Vendor API keys for Groq, Gemini, OpenAI, Anthropic, and xAI are written only by an administrator, through the providers settings route. They are encrypted at rest (smenc:v1) with INSTANCE_ENCRYPTION_KEY or a local instance key file. Responses return a masked last-four hint. Azure chat is not wired. xAI chat is wired (src/lib/models/xai.ts) and its key can be stored, but it has not been exercised end to end, so this guide does not treat it as a tested path.")
    body(doc, "AWS access keys are never accepted in the UI and never stored in the database. A pasted AWS key is rejected with 422 at the route and again when instance configuration is saved. Bedrock uses the AWS SDK credential chain: environment variables, a shared profile or SSO, a container role, or an instance role. The awsProfile setting is only a pointer into the shared config file. The profile list reads section headers, not secret values. ALLOW_REAL_AWS=false is a hard kill switch. The administrator's toggle cannot override it. canUseBedrock() is the single gate. Billable AWS calls stay off until that gate is open.")
    body(doc, "Any AWS SDK client other than Cognito is constructed inside withRealAwsEndpoints(). A global endpoint variable, which the local AWS stand-in sets, would otherwise be picked up by every client, break SSO credential resolution, and send Bedrock traffic to localhost. Cognito reads its endpoint explicitly and is left alone.")

    body(doc, "A task model is a small helper for one narrow job, separate from chat routing (src/lib/models/task-models.ts). The one job today is formatting a reply into a downloadable file when the chat model did not follow the file format. It is off until an administrator chooses a model under Settings, Models. It uses exactly that model, with the provider locked, and never falls back to another one. An outside vendor needs the vendor consent notice accepted. It sees only the reply text with the non-exemptible categories already replaced. The setting is not a secret and is stored like model availability.")

    h1(doc, "11  Prompts, evidence, and citations")
    body(doc, "buildConversationalBase() applies on every turn: natural, concise, and not forced into an operations pitch. Operational instructions and the workspace context are appended only when the turn policy says the turn needs them. The workspace name is the speaker the user sees. It is not pasted into a social prompt as a persona.")
    body(doc, "Retrieved text is treated as evidence, not as instructions. src/lib/prompt-guard.ts neutralizes injection patterns in that text and fences it so a passage inside a document cannot retask the model as if it were the system prompt. Conversation history is included as continuity. It is not a Source, and the model is not told to cite a prior exchange or its own general knowledge as a document.")
    body(doc, "When Sources are shown, they are the passages that were packed for that turn and that the answer actually cited. After the answer, auditCitations in src/lib/citation-audit.ts checks the [n] markers. A marker that does not match a packed index is invalid. Packed sources the answer never cites are unused. An answer that admits it does not have the evidence, and cites nothing, is recorded as declined rather than as a successful citation. The audit does not use a second model to judge whether a claim is true.")
    body(doc, "Ranking uses a fused score (src/lib/retrieval-rank.ts): cosine similarity against the workspace's embeddings plus a small keyword bonus. Candidates must clear a vector floor of 0.30 before scoring. A grounded turn accepts a fused score of 0.40, and a silent search on a general question needs 0.55. There is no single fixed similarity floor, because one dropped correct passages. The search is limited to the current user's documents in the current workspace, and to the document category. Fact rows are not in that index.")

    h2(doc, "Files from chat")
    body(doc, "detectFileRequest (src/lib/turn-policy/file-request.ts) recognises a request such as export this as Excel or give me a Word document. It changes only the shape of the reply: fileOutputInstructions is appended for that turn and the reply limit rises to FILE_REPLY_TOKENS. It never changes retrieval, Sources or the tier. The model writes the file in one fenced block whose info string is file:name.ext. The browser builds the file when the user presses Download (src/lib/files/): xlsx or csv from CSV, with several sheets separated by a sheet marker line, docx or md from Markdown, and a few other allowlisted types as text. Nothing is stored on the server. File names are cleaned and allowlisted, and anything that would run when opened is saved as .txt. A spreadsheet cell that looks like a formula is written as text. A file over 2 MB is refused.")
    body(doc, "If a reply to a file request has no usable block, applyFileFallback (src/lib/files/fallback.ts) first rewrites a CSV block, a Markdown table or a Markdown document into the same block, locally and for free. Only when that finds nothing, and an administrator has approved a task model, does POST /api/files/format ask that one model to build the file.")

    h1(doc, "12  Knowledge")
    body(doc, "A file added in Knowledge, or attached in chat, belongs to the current workspace. The server extracts text, scrubs it with the instance detection policy, chunks the scrubbed text, and embeds the chunks. On a local installation the original file is also stored on disk, outside Postgres. The hosted path does not use that local file step. Converting an image into a vector does not make the image safe. The extracted text is what gets scrubbed, then embedded.")
    body(doc, "Chunks are about 1,500 characters, with about 200 characters of overlap, and pieces shorter than 80 characters are dropped. Breaks prefer a paragraph or a sentence boundary. Local embeddings use nomic-embed-text through Ollama, 768 dimensions. Hosted embeddings target Amazon Titan v2, 1,024 dimensions, and that dimension is fixed when the DynamoDB vector index is created. The two stores are not mixed in one installation.")
    body(doc, "Spreadsheets can be queried through a constrained analyze route rather than by handing the model the whole grid. Images are described with Groq vision when a key exists, otherwise the ingest keeps a filename stub, and that text is scrubbed like any other extract. Reprocess retries extract and embed from the stored original on the local path.")
    body(doc, "Uploads are untrusted, so limits apply before a parser runs (src/lib/documents/upload-guard.ts). A file is at most 10 MB. A Word or Excel file is a zip container. Its directory is checked first: more than 100 MB unpacked, an expansion above 100 times, more than 5,000 parts, or a ZIP64 archive is refused. Those sizes are what the archive claims about itself, so each part is then inflated with a hard output cap. A part that unpacks larger than it declared, or past the remaining budget, is refused before a document parser sees it, and only stored or deflate compression is accepted. Extraction has a 60-second limit. The size of the whole request is read from its header before the body is parsed, so an oversized upload is refused early. Each person can add 200 files a day and store 1 GB in total. The limits can be raised with environment variables on a larger installation.")

    h1(doc, "13  Memory and bookmarks")
    body(doc, "Memory has two layers, and they are not the same store. Working context is recent messages plus a session summary. It keeps a conversation continuous. It is not a citation. Long-term facts are created from what the user wrote, through the summarize route. Assistant replies are not stored as facts and are not embedded as memory fragments.")
    body(doc, "The user can correct a fact, pin it, or forget it. Forget marks it forgotten and keeps the row so it can be restored. Delete removes the fact. The vector index used for citations remains document chunks only.")
    body(doc, "A bookmark marks one message. Choosing one scrolls to that message and highlights it. Removing it drops the bookmark. A bookmark does not become a fact and does not become a Knowledge document. Emoji that appear on a bookmark row are taken from the preview so the row is recognizable. They are not a third memory system.")
    body(doc, "Emoji shows up in three other places, and they should not be collapsed. The workspace mark is identity in the sidebar. The composer control inserts a character into the draft as ordinary text. Profile avatars are a separate set: a preset glyph, or a cropped JPEG or PNG stored on local disk or in a private bucket and served only to a signed-in session through the avatar route. There is no public avatar URL.")

    h1(doc, "14  Voice")
    body(doc, "Voice is off until the operator selects a backend. Two backends exist: Groq Whisper, and Amazon Transcribe. Audio for a hosted transcription is sent to that service before any text detection can run. The transcript is inserted into the composer. Detection runs when the user sends that text, the same as typed text. The waveform is not stored as the chat message.")
    body(doc, "Whisper posts audio to Groq and can preview a growing recording in chunks while the microphone is open. Transcribe is the streaming path on hosted Chromium. When the hosted install has its signing role (TRANSCRIBE_STREAM_ROLE_ARN), the browser streams the microphone straight to Transcribe with a five-minute signed address, so words appear while the user speaks; the role can only open a Transcribe stream, and a session ends after ten minutes. It is billed on the operator's AWS account, behind the same real-AWS gate as Bedrock, and counts toward the monthly cost cap. If live captions fail, the whole recording is transcribed once instead. A choice named whisper that is not one of these two backends is ignored.")
    body(doc, "Why the browser connects directly: Amplify's server receives an upload in one piece, so streaming the microphone through a SecureMind route could not show words while the user speaks, and API Gateway does not stream request bodies either. POST /api/transcribe/live-session therefore checks identity, the hosted edition, the Admin's choice and audio acknowledgement, the cost cap and a per-user rate limit, then assumes a separate one-purpose role (securemind-transcribe-stream) that can only open a Transcribe stream, with a session policy of the same scope. It signs a WebSocket address with those temporary credentials using AWS Signature Version 4 (src/lib/transcribe/presign.ts). The address is good for five minutes, is never logged, and never contains the secret key. The page's Content-Security-Policy allows exactly one extra connection origin, the Transcribe streaming endpoint of the deployment's region, and only when this feature is configured. The browser sends audio in the AWS event-stream format (src/lib/transcribe/event-stream.ts, src/lib/transcribe/direct-client.ts). The decision is recorded in ADR 0009. Without the role and region settings, the older route through the server still works.")

    h1(doc, "15  Cost estimates and caps")
    body(doc, "Every inference event can store an estimate. The estimate is recorded input tokens and output tokens, times a rate, then adjusted so the figure tends high. It is not read from AWS Cost Explorer, and it does not include hardware, seats, or other infrastructure. The provider's invoice remains the charge.")
    h2(doc, "The formula")
    body(doc, "Rates are USD per million tokens. The raw amount is (input tokens times the input rate, plus output tokens times the output rate) divided by 1,000,000. That function is costFromRates in src/lib/pricing/types.ts.")
    body(doc, "estimateInferenceCost then adjusts the raw amount. Ollama is $0 and is marked as priced, so a local call does not show up as an unknown price. For every other provider the order is: an administrator rate override if one exists for that provider and model; otherwise the Groq table or the Bedrock on-demand table; otherwise a floor.")
    body(doc, "The floor is the Max-tier Groq list rate, llama-3.3-70b-versatile at 0.59 and 0.79 per million. An unknown cloud model, including a Groq or Bedrock id that is missing from the table, uses that floor. The floor exists so a missing price cannot hide spend or slip past a cap.")
    body(doc, "The raw amount, or the floor, is then multiplied by 1.10. The result is rounded up to the next 0.0001 dollar. A zero or non-finite raw amount stays zero. The constants are CONSERVATIVE_COST_UPLIFT and COST_ROUND_UP_QUANTUM in src/lib/pricing/index.ts. An override is stamped as an override and still receives the uplift. Overrides apply to new turns only. They do not rewrite events already stored.")
    h2(doc, "A worked example")
    body(doc, "A Standard turn on openai/gpt-oss-120b uses 1,000 input tokens and 500 output tokens. The raw amount is (1,000 times 0.15 plus 500 times 0.60) divided by 1,000,000, which is 0.00045 dollar. Times 1.10 is 0.000495. Rounded up to the next 0.0001 is 0.0005 dollar. That is the number stored and the number the cap adds.")
    body(doc, "The same token counts priced as the savings baseline, llama-3.3-70b-versatile, are 0.000985 raw, 0.0010835 after the uplift, and 0.0011 after the round-up. Savings are the baseline minus the actual estimate, and never below zero: 0.0006 dollar in this example. The baseline model is fixed in code. It does not follow a different model the operator may have assigned to Max. Savings illustrate the difference from always using that list-price model. They are not a refund. A local Ollama reply is $0, so its savings figure is the entire uplifted baseline for those tokens.")
    h2(doc, "The monthly cap")
    body(doc, "The cap is a per-user sum of estimatedCostUsd for the current UTC calendar month. The month starts at 00:00 UTC on the first day. The default cap is 10 dollars. Null means unlimited. A local owner with no profile row is unlimited until a cap is set. Other users without a profile receive the instance default. The maximum value the settings route accepts is 100,000 dollars, with at most two decimal places.")
    body(doc, "Warnings are ratios of spend to cap: 50 percent and 80 percent. At 100 percent the status is blocked. Before a chat call, if the user is already at the cap and the call is billable, the server returns HTTP 402 and does not invoke the model. Billable providers are Groq, Bedrock, OpenAI, Anthropic, Google, and xAI. An image counts as billable even before the provider check. Voice is checked at its own routes. Ollama is outside the cap, so local models can still run after a cloud cap is reached. Raising the cap is an administrator action. The block cannot make the provider's bill equal the estimate.")

    h2(doc, "Voice")
    body(doc, "Amazon Transcribe is billed per second of audio on the operator's account. SecureMind records its use as an inference event with the provider transcribe, so it counts toward the same monthly cap and shows on the Cost page. The list rate used is 0.024 dollar a minute (TRANSCRIBE_USD_PER_MINUTE in src/lib/transcribe/voice-usage.ts). A recording sent through the server is recorded at its real length. A direct live session runs from the browser to AWS, so the server never learns how long it lasted. It is recorded as one flat minute when its address is signed, and the browser ends it after ten minutes. A modified browser could keep a connection open longer, so the invite-only sign-in, the per-user rate limit and the monthly budget alarm are the backstops. Voice is not a chat reply, so the Usage page does not count it.")

    h1(doc, "16  Insights")
    body(doc, "Insights is the right-hand panel: Models, Usage, Cost, and Security. Each opens a full page. It follows the workspace selector and the range Today, Weekly, Monthly, or All time. The counts are for the signed-in account. An administrator can open another person's Usage, Cost, and Security from the directory. The views do not contain raw prompts or raw detected values.")
    body(doc, "The blue number on a tier is the count of replies routed to that tier in the selected workspace and range. It is not a count of calls to one model id. Provider chips are how many models of that provider can be assigned. They are not call counts. Hide does not remove a model from a tier that is already pointing at it.")
    body(doc, "Security groups findings by category, source, and outcome, with a coarse risk level derived from the category. It reports what the scanner flagged. It is not a vulnerability scanner and not a score of data leaving the system. Usage stores which model and tier ran, token counts, grounding status, and success or failure. Cost uses the estimates from section 15, split by time, model, workspace, and input versus output.")

    h1(doc, "17  People, profile, and appearance")
    body(doc, "Profile stores a display name, a description, and an avatar for the signed-in user. Account security holds password change and optional two-step verification, and forgot-password uses Cognito when that is the sign-in service. Forgot-password responses are written so they do not reveal whether the email exists. Theme is light or dark, stored on the device under the key sm_theme. Dark is the default until a choice is saved. The choice does not follow the user to another browser.")
    body(doc, "The status chips above the composer report detection, document parsing, routing, and voice. They are visible until the user collapses them. That collapse is stored on the device. It is not an account setting.")
    body(doc, "The administrator's directory shows email, role, month-to-date estimate against the cap, and sign-in status. Centralized Insights on that screen are installation totals. Per-user drill-down reuses the same Usage, Cost, and Security views, scoped to that member.")
    body(doc, "A member's first workspace opens with a short welcome that says where things are, not the connect-a-model message an owner sees, and a Take the tour button that stops showing after three visits or when it is closed (src/lib/client/tour-button.ts). The Admin page also has a Portal card that sets the logo, name and short phrase at the top-left of the workspace; the sign-in page, the browser title and email stay SecureMind.")
    body(doc, "An Admin can set a Member's password, which lets that Admin sign in as the Member and see their conversations. That is normal inside one team, and it is why separate organisations should not share an installation. Keep the number of Admins small and have them turn on two-step verification.")

    h1(doc, "18  Where data is stored")
    body(doc, "Local path. The browser talks to Next.js on port 3000. Postgres holds workspaces, messages, documents, facts, profiles, and analytics. pgvector holds 768-dimension embeddings. The vector column is created outside Prisma and queried with raw SQL that must filter by user id. Sign-in in Docker goes to Floci. Chat and embeddings can go to Ollama on port 11434. Uploads, avatars, and the instance encryption key are files under the data directory. Optional Bedrock uses one AWS profile folder the operator chooses to mount (docker-compose.bedrock.yml), never the whole ~/.aws directory. Optional vendor calls use the stored key.")
    body(doc, "Docker exposure. Only the app is published, on 127.0.0.1:3000. Postgres, Floci and Ollama are reachable only inside the Compose network. To reach the app from other computers, the operator adds the included HTTPS front end (docker-compose.tls.yml, Caddy), which also turns on HSTS. Containers run with no-new-privileges, process and memory limits and rotated logs. The app runs as a non-root user with no Linux capabilities and a read-only filesystem apart from /data and scratch paths. Images are pinned by digest, the app runs on Node.js 22, and the shipped image contains no source files or test fixtures. Customers build the image from source.")
    body(doc, "Hosted path. Amplify serves the UI, sign-in, workspaces, settings, and the non-chat APIs. The chat stream goes to API Gateway and the chat Lambda. Application records and the vector index use DynamoDB. Cognito is the user pool. Systems Manager Parameter Store holds the hosted secrets as encrypted SecureString parameters: the instance encryption key and first-run setup code, the vendor keys an administrator saves, and an optional bootstrap Groq key. Standard parameters have no monthly charge. The table name is a plain setting, not a secret. Profile photos use a private bucket. CloudWatch holds Lambda logs. Prompt and response content logging is not turned on for that dashboard. The Amplify server runs on its own compute role, separate from the build role, with only the runtime permissions it needs. DynamoDB has point-in-time recovery and deletion protection; the Cognito pool has deletion protection on by default (a sandbox that is torn down sets it to INACTIVE). A CloudTrail trail records account activity to a private, encrypted, versioned bucket. Five CloudWatch alarms (chat errors and throttles, API 5xx, an API 4xx spike, DynamoDB throttles) and a monthly AWS Budget notify the email set as ALERT_EMAIL. The chat Lambda runs with NODE_ENV set to production, and the hosted API accepts only Cognito sign-in. The Amplify server logs expire after 30 days. Amplify's server role may assume one more role, securemind-transcribe-stream, which can only open a Transcribe stream. There is no custom VPC in the current shape. The region is chosen in the infrastructure configuration. The in-app setup wizard asks for a storage preference and a risk acknowledgement. It does not ask for a region.")
    body(doc, "A push to the application main branch runs tests, a type check and a production build. Deploying from GitHub Actions is opt-in: only when the repository variable AWS_DEPLOY_ENABLED is true and the role and state bucket are configured does it apply Terraform and start an Amplify release. Pull requests plan only, and they use a separate read-only role (securemind-gha-plan) that cannot read data, never the deploy role. GitHub Actions are pinned by commit, and Dependabot proposes updates. Neon and Aurora are not the current database. They are not mixed with DynamoDB.")

    h1(doc, "19  Operating an installation")
    body(doc, "Team size. One installation is for one organisation. It is recommended for 1 to 25 people and can serve up to about 100 people with the adjustments below. Beyond that, or for several unrelated organisations, run separate installations.")
    add_table(
        doc,
        ["To support", "What to do"],
        [
            ["More than about 10 people chatting at once (AWS)", "Ask AWS to raise Lambda concurrent executions (Service Quotas, free). New accounts often start at 10, and each streaming reply holds one slot for up to five minutes."],
            ["Larger teams in general", "Keep two or three Admins and have them turn on two-step verification. Review the audit log."],
            ["Spend", "Each person's default cap is 10 dollars a month, so 100 heavy users could reach about 1,000 dollars in model costs. Set caps and the monthly budget to match."],
            ["Docker beyond a few people", "Run it on a server, add a model API key (the local model serves only a few people), and use the HTTPS front end and a firewall that allows ports 80 and 443 only."],
        ],
        [2.4, 4.3],
    )
    body(doc, "Monitoring. On AWS, set ALERT_EMAIL (and optionally MONTHLY_BUDGET_USD) in .env.deploy before applying, then confirm the email AWS sends. COGNITO_DELETION_PROTECTION=ACTIVE belongs in the same file for a real installation. These values stay in that git-ignored file and in the operator's own account.")
    body(doc, "Web application firewall. AWS WAF is not included, to keep an idle installation near zero cost. The gateway's Cognito check, its request throttle, the 4xx alarm and the budget cover the main risk. An operator who wants WAF can add it to their own account.")

    h1(doc, "20  What this system does not do")
    body(doc, "It does not judge answers with a second model. It does not scan the model's output for sensitive data. It does not re-scrub old chat when the detection tier changes. It does not restore a redacted value. It does not offer a public avatar URL, a hard purge of a user's data, a per-workspace provider policy, or Bedrock Guardrails as a product feature. It cannot cap the length of a browser-to-AWS voice stream on the server, only estimate it. It does not revoke an access token that has already been issued before it expires.")
    body(doc, "It does not accept a pasted AWS access key. It does not ship a product-owned Groq key as a silent default on a hosted install. It does not query the provider's billing API to draw the cost charts. It does not treat a region setting as a promise that every connected service processes data only in that region.")
    body(doc, "KMS, Macie, and a production multi-tenant auth system are not features of this application. One installation is not a boundary between separate organisations, and the Docker sign-in (Floci) is not meant to face the public internet. An operator who needs those controls provides them outside this product.")

    h1(doc, "21  Where the rules live")
    add_table(
        doc,
        ["Rule", "Where"],
        [
            ["Turn kind, retrieval, Sources", "src/lib/turn-policy/decide-turn.ts"],
            ["Tier score, carry-over, ambiguous band", "src/lib/models/route-message.ts"],
            ["Default models for each tier", "src/lib/models/router.ts"],
            ["Classifier prompt", "src/lib/models/groq.ts"],
            ["Allowlist", "src/lib/models/allowlist.ts"],
            ["Chat orchestration, vision override, 402", "src/lib/chat/run-chat-stream.ts"],
            ["Detection families and safe terms", "src/lib/pii.ts"],
            ["Document scrub, chunk sizes", "src/lib/documents.ts"],
            ["Fused retrieval score and thresholds", "src/lib/retrieval-rank.ts"],
            ["Citation check", "src/lib/citation-audit.ts"],
            ["Evidence fencing", "src/lib/prompt-guard.ts"],
            ["Rate tables and uplift", "src/lib/pricing/index.ts"],
            ["Monthly cap", "src/lib/cost-cap.ts"],
            ["Session user id", "src/lib/identity.ts"],
            ["Token claims, fail-closed defaults", "src/lib/auth-claims.ts"],
            ["Server-side chat check", "src/lib/chat/server-scrub.ts"],
            ["Sign-in rate limits", "src/lib/rate-limit.ts"],
            ["Upload limits and the unzip cap", "src/lib/documents/upload-guard.ts, src/lib/documents/zip-guard.ts"],
            ["Score card and evidence floors", "src/lib/models/route-scorecard.ts, applyEvidenceFloors in src/lib/models/route-message.ts"],
            ["Tier meanings, routing fixtures", "src/lib/models/tier-definitions.ts, src/lib/models/route-fixtures.ts"],
            ["Retired-model unlisting", "src/lib/models/retire-failed-model.ts"],
            ["Task models and file building", "src/lib/models/task-models.ts, src/lib/turn-policy/file-request.ts, src/lib/files/"],
            ["Voice cost and cap", "src/lib/transcribe/voice-usage.ts"],
            ["Direct live captions", "src/lib/transcribe/stream-session.ts, presign.ts, event-stream.ts, direct-client.ts"],
            ["Server check outside chat", "scrubServerText in src/lib/chat/server-scrub.ts"],
            ["Admin audit log", "src/lib/audit-log.ts"],
            ["AWS credential chain", "src/lib/aws-credentials.ts"],
        ],
        [3.2, 3.5],
    )

    try:
        doc.save(OUT)
        print(OUT)
    except PermissionError:
        alt = OUT.with_name("SecureMind-Technical-Guide-v1.docx")
        doc.save(alt)
        print(alt)


if __name__ == "__main__":
    build()
