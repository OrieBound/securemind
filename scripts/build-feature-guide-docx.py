"""Build the SecureMind feature guide Word document.

Regenerate after edits:

    python scripts/build-feature-guide-docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor, Twips

OUT = Path(__file__).resolve().parents[1] / "docs" / "SecureMind-Feature-Guide.docx"

# Light theme of the application (src/app/globals.css :root).
INK = RGBColor(0x1A, 0x1C, 0x1F)
MUTED = RGBColor(0x5C, 0x61, 0x6B)
VIOLET = RGBColor(0x5B, 0x4D, 0xB8)
VIOLET_INK = RGBColor(0x3A, 0x32, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x17, 0x6B, 0x42)
AMBER = RGBColor(0x8A, 0x5A, 0x00)

PRIMARY = "5B4DB8"
SOFT = "E4E2F2"
CANVAS = "E6E8ED"
CARD = "F4F5F7"
SIDEBAR = "DDE0E6"
RULE = "D0D3DA"
SUCCESS_SOFT = "DDECE4"
WARNING_SOFT = "F3E8C8"


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run_font(run, name: str, size: int, color: RGBColor, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_text(paragraph, text: str, *, size=11, color=INK, bold=False, italic=False, name="Calibri") -> None:
    run = paragraph.add_run(text)
    set_run_font(run, name, size, color, bold)
    run.italic = italic


def body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text)


def purpose(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, "Purpose. ", size=11, color=VIOLET, bold=True)
    add_text(p, text)


def h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        set_run_font(run, "Calibri", 18, VIOLET, True)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)


def h2(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        set_run_font(run, "Calibri", 13, VIOLET_INK, True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def set_cell_border(cell, color: str = RULE) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def set_cell_margins(cell, twips: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def write_cell(cell, text: str, *, header: bool = False, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    color = WHITE if header else INK
    add_text(p, text, size=10, color=color, bold=header or bold)
    set_cell_border(cell)
    set_cell_margins(cell)
    if header:
        shade(cell, PRIMARY)
    elif fill:
        shade(cell, fill)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = False
    table.allow_autofit = False
    for i, heading in enumerate(headers):
        write_cell(table.rows[0].cells[i], heading, header=True)
        table.rows[0].cells[i].width = Inches(widths[i])
    for r, row in enumerate(rows):
        fill = CARD if r % 2 == 0 else "EEF0F3"
        for c, value in enumerate(row):
            write_cell(table.rows[r + 1].cells[c], value, bold=(c == 0), fill=fill)
            table.rows[r + 1].cells[c].width = Inches(widths[c])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    r = run._r

    def fld(kind: str) -> None:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r.append(el)

    fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    r.append(instr)
    fld("separate")
    text = OxmlElement("w:t")
    text.text = " "
    r.append(text)
    fld("end")


def header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    p_pr = hp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "right")
    tab_el.set(qn("w:pos"), "9360")
    tabs.append(tab_el)
    p_pr.append(tabs)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), PRIMARY)
    border.append(bottom)
    p_pr.append(border)
    hp.paragraph_format.space_after = Pt(4)
    add_text(hp, "SECUREMIND", size=9, color=VIOLET, bold=True)
    add_text(hp, "   ·   Feature guide", size=9, color=MUTED)
    hp.add_run("\t")
    add_text(hp, "Version  1.2", size=9, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    p_pr = fp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "right")
    tab_el.set(qn("w:pos"), "9360")
    tabs.append(tab_el)
    p_pr.append(tabs)
    top = OxmlElement("w:pBdr")
    top_el = OxmlElement("w:top")
    top_el.set(qn("w:val"), "single")
    top_el.set(qn("w:sz"), "6")
    top_el.set(qn("w:space"), "1")
    top_el.set(qn("w:color"), RULE)
    top.append(top_el)
    p_pr.append(top)
    add_text(fp, "Version 1.2", size=9, color=MUTED)
    fp.add_run("\t")
    add_text(fp, "Page ", size=9, color=MUTED)
    add_field(fp, " PAGE ")
    add_text(fp, " of ", size=9, color=MUTED)
    add_field(fp, " NUMPAGES ")


def figure(doc: Document, number: str, caption: str) -> None:
    """A caption in the text. Paste the picture on this line; it is not boxed."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.08
    add_text(p, f"Figure {number}. ", size=10, color=VIOLET, bold=True)
    add_text(p, caption, size=10, color=MUTED)


def swatch_row(table, row_index: int, fill: str, name: str, meaning: str) -> None:
    color_cell = table.rows[row_index].cells[0]
    color_cell.text = ""
    p = color_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, name, size=9, color=WHITE if fill in (PRIMARY, "1A1C1F") else INK, bold=True)
    shade(color_cell, fill)
    set_cell_border(color_cell, RULE)
    set_cell_margins(color_cell, 80)
    write_cell(table.rows[row_index].cells[1], meaning, fill=CARD)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.95)
    section.bottom_margin = Inches(0.85)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # Cover band
    banner = doc.add_table(rows=1, cols=1)
    cell = banner.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "SECUREMIND", size=11, color=WHITE, bold=True)
    t = cell.add_paragraph()
    t.paragraph_format.space_before = Pt(8)
    t.paragraph_format.space_after = Pt(2)
    add_text(t, "Feature Guide", size=28, color=WHITE, bold=True)
    s = cell.add_paragraph()
    s.paragraph_format.space_before = Pt(2)
    s.paragraph_format.space_after = Pt(6)
    add_text(s, "Application overview and feature purposes", size=12, color=RGBColor(0xE4, 0xE2, 0xF2))
    shade(cell, PRIMARY)
    set_cell_border(cell, PRIMARY)
    set_cell_margins(cell, 160)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(10)
    meta.paragraph_format.space_after = Pt(8)
    add_text(meta, "Version 1.2", size=11, color=VIOLET, bold=True)
    add_text(meta, "   ·   25 October 2026", size=11, color=MUTED)

    body(
        doc,
        "A guide for people who want to understand SecureMind, choose how they use AI, and see how their conversations, models, costs, and sensitive-data findings fit together.",
    )
    body(
        doc,
        "For individual users, project owners, developers, and administrators running their own installation.",
    )

    h1(doc, "1  Purpose and vision")
    body(
        doc,
        "SecureMind brings conversations, documents, and useful context into organized workspaces. Its purpose is to help people use AI with more control over model choice, sensitive information, and estimated cost.",
    )
    body(
        doc,
        "The vision is a personal working space that remains useful as a project develops. It is especially suited to developers and people working on their own projects who want an additional layer of reassurance and transparency around their use of AI.",
    )
    body(
        doc,
        "Model tiers let users reserve more expensive models for more complex questions, while using other models for simpler requests. Memory and Knowledge provide continuity, and Insights makes model usage, estimated costs, and detected sensitive-data concerns visible.",
    )
    body(
        doc,
        "Accessibility also means being able to run your own installation: through Docker on your own machine, or in a cloud environment you configure. Each installation has its own users, settings, and provider connections. SecureMind is not intended to operate one public service holding everyone’s API keys.",
    )
    body(
        doc,
        "An administrator manages provider keys, access, and shared settings. Members use that setup and can have individual estimated spending caps. Stored application data is held within the installation’s environment. Configured external providers can still receive content needed for model responses, transcription, or other processing.",
    )
    body(
        doc,
        "SecureMind is intended for small to mid-sized teams. One installation serves one organization: 1 to 25 people is the recommended range, and it can serve up to about 100 people with a few adjustments, mainly an AWS concurrency limit and a realistic budget. The technical guide lists them. Separate organizations should each run their own installation.",
    )

    h1(doc, "2  How to use this guide")
    body(
        doc,
        "Each section explains what a feature does and why it exists. Start with Workspaces and Chat for everyday use, then explore Knowledge, Memory, Bookmarks, and Insights. Settings and the Admin console explain how the installation is configured and managed.",
    )
    body(
        doc,
        "The Architecture and Technical Guide covers implementation details: processing order, storage, access boundaries, authentication, cost calculations, and the algorithms behind complexity routing.",
    )
    body(doc, "Guide sections")
    body(
        doc,
        "3 Terms and controls  ·  4 Connections and credentials  ·  5 Installation options  ·  6 Workspaces  ·  7 Chat  ·  8 Knowledge  ·  9 Memory  ·  10 Bookmarks  ·  11 Insights  ·  12 Model settings  ·  13 Detection policy  ·  14 Profile and account security  ·  15 Appearance  ·  16 Admin console  ·  17 Getting started  ·  18 How the parts fit  ·  19 Scope and limitations",
    )

    h1(doc, "3  Terms and controls")
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Installation", "Your own running copy of SecureMind, with its users, settings, data, and provider connections."],
            ["Administrator", "Manages accounts, roles, provider connections, detection settings, and estimated spending caps."],
            ["Member", "Uses the connections and shared settings established by the administrator, without viewing or changing secret keys."],
            ["Workspace", "Organizes one body of work: its chat, Knowledge, memory, and bookmarks."],
            ["Complexity tiers", "Max, Plus, Standard, Light, and Micro. Assign a model to each tier. Routing selects a tier for the request."],
            ["Fixed model", "Use one selected model for replies instead of complexity routing."],
            ["Enable", "Checks one model that is not in the chat list yet. If it answers, it joins the list."],
            ["Test", "Runs that same check on a model that is already in the chat list."],
            ["Check access", "Runs the check once for each listed model on a provider key."],
            ["Hide and Show", "Hide takes a model off the chat list. Show puts it back. A hidden model can still be the model a tier runs."],
            ["Check failed", "The last check did not succeed. The label does not explain the cause. Hidden is a separate state."],
            ["Provider chip", "Shows how many models of that provider, such as Gemini or Anthropic, can be assigned. It is not a request count."],
            ["Blue count", "How many replies were routed to that tier, for the workspace and date range selected in Insights."],
            ["Vendor API key", "A provider credential managed by the administrator. Groq and xAI, the provider of Grok, are separate providers."],
            ["Score card", "The \"Why this tier?\" link under a reply. It lists the signals that raised or lowered the score, any adjustment, and the tier that ran."],
            ["Task model", "A helper model an administrator approves for one narrow job, such as turning a reply into a file. It is off until chosen."],
            ["Portal", "The logo, name, and short phrase in the top-left of the workspace, set by an administrator."],
            ["Cost estimate", "An approximate usage cost based on recorded activity and configured provider rates, not an invoice."],
            ["Detection", "Best-effort identification of sensitive patterns so matched content can be reviewed and redacted."],
        ],
        [1.7, 4.9],
    )

    h1(doc, "4  Connections and credentials")
    purpose(doc, "Connect the model and transcription services available to the installation, with credentials controlled by an administrator.")
    body(
        doc,
        "The administrator adds provider API keys and configures shared connections. Members use those connections without seeing the secret keys or changing them. Vendor keys are stored encrypted. Settings shows a masked hint, not the secret.",
    )
    body(
        doc,
        "Provider connections include Groq, Gemini, OpenAI, Anthropic, and xAI for Grok. Groq and xAI are separate providers.",
    )
    body(
        doc,
        "Amazon Bedrock uses AWS authentication configured for the installation. On a local installation, Bedrock uses the credential chain already on that computer: environment variables, a shared profile or SSO, a container role, or an instance role. SecureMind does not accept pasted AWS access keys. On a cloud installation, Bedrock uses that deployment’s IAM role.",
    )
    body(
        doc,
        "Local Ollama models run through Docker in the local installation. The first boot pulls nomic-embed-text for Knowledge embeddings and llama3.2:3b for local chat. They do not need a vendor API key, but a vendor API key (or Bedrock on AWS) is strongly recommended. The local models use the computer's own resources, and unless the machine is very powerful they are slow and will not do well. Their speed depends on available hardware and may be substantially slower than hosted alternatives. Ollama is not part of the cloud configuration.",
    )
    h2(doc, "Voice transcription")
    body(
        doc,
        "Voice input uses a configured transcription service: Amazon Transcribe or Groq Whisper. The transcript appears in the composer for review before sending. For hosted transcription, audio is sent to the selected service before the resulting text can be scanned. On the hosted edition with Amazon Transcribe, words appear while you speak, a session ends after ten minutes, and voice counts toward the monthly cost cap and is refused once a person is over it. Your browser connects to Amazon Transcribe with a short-lived address that the server signs, so the audio does not pass through the SecureMind server. If live captions fail for any reason, the whole recording is transcribed in one pass when you press the tick.",
    )
    h2(doc, "Task models")
    body(
        doc,
        "Chat chooses a model by complexity tier. A task model is separate: a helper an administrator approves for one narrow job. The one job today is building a downloadable file when the chat model did not produce the file format. It is off until an administrator chooses a model under Settings, Models. It always uses exactly the model that was chosen, and never falls back to another provider. It sees only the reply text, after card numbers, social security numbers and secrets have been replaced, and a vendor model needs the vendor notice accepted first.",
    )

    h1(doc, "5  Installation options")
    purpose(doc, "Let people operate SecureMind in an environment they control.")
    h2(doc, "On your own machine")
    body(
        doc,
        "The local option uses Docker and includes Ollama, as described above. Users can configure external providers when they want hosted models or transcription. Running the application locally does not make external provider calls local.",
    )
    body(
        doc,
        "Docker builds SecureMind from the source you download, and by default it is reachable only from that computer. The first account needs a one-time setup code, which the installation prints when it starts. To use it from other computers on your network, add the included HTTPS front end and a firewall. Docker is meant for a computer or a private network, not for the public internet.",
    )
    h2(doc, "In your cloud environment")
    body(
        doc,
        "The cloud option runs in the administrator’s own AWS environment. The region is set in the infrastructure configuration. The in-app setup wizard does not ask for a region. Ollama is not part of SecureMind’s cloud configuration. Provider availability and any cross-region processing depend on the services and models configured.",
    )
    body(
        doc,
        "The cloud option includes backups of the data (point-in-time recovery), protection against deleting the database or the user directory by mistake, an AWS activity trail, alarms, and a monthly budget alert sent to an email the administrator chooses. Sign-in is invite-only and is checked again at the API before any chat runs.",
    )
    body(
        doc,
        "Setup for a local Docker installation and for a deployment in your own AWS account is on the SecureMind website, and linked from the GitHub README.",
    )

    h1(doc, "6  Workspaces")
    purpose(doc, "Keep the conversations and context for one project or body of work together.")
    body(
        doc,
        "Each workspace contains its own chat, Knowledge, memory, and saved bookmarks. Retrieval stays within the selected workspace. Its name, emoji, and color help you recognize it. They do not select a model or define an AI personality.",
    )
    figure(doc, "1", "The workspace list. Name, emoji, and color identify it. They do not choose a model.")
    body(
        doc,
        "Switching workspaces opens that workspace’s conversation and context. An unsent queued message is discarded, and the composer is cleared, so the message cannot be delivered into the wrong workspace.",
    )

    h1(doc, "7  Chat")
    purpose(doc, "Ask questions, review sensitive text, and receive replies using your chosen model configuration.")
    h2(doc, "Review before processing")
    body(
        doc,
        "In the typed-message flow, the draft remains in the browser while detection runs. Sensitive details the review finds are redacted, and you see them before the message is processed. The reviewed text is what continues.",
    )
    body(
        doc,
        "The normal path does not send the original detected value. Detection is best effort: a value that is not detected may remain in the submitted text. A request that skips the browser skips this review, so the server checks every chat message again for the highest-risk values (social security and social insurance numbers, card numbers and secrets) and replaces them. Other details are reviewed in the browser only. The scan is risk reduction, not a guarantee.",
    )
    h2(doc, "Choosing the answering model")
    body(
        doc,
        "Complexity routing selects a tier and calls the model assigned to it. You can assign more expensive models to higher-complexity tiers and use other models for simpler requests. Fixed model uses the single model you select.",
    )
    body(
        doc,
        "Every reply carries a \"Why this tier?\" link. It opens a score card: the baseline, each signal that fired with its points, any adjustment, and the tier that ran. No message text is stored on the card. A message that is answered from your documents is never run on the cheapest tiers: a short question about a document runs on at least Standard, and a summary or analysis across documents runs on at least Plus.",
    )
    body(
        doc,
        "If a provider retires a model, the next reply that picks it fails with the provider's answer. SecureMind then takes that model off the chat menus, says so in the reply, and leaves it off until a later Test passes. A temporary problem such as a rate limit or an outage does not unlist anything.",
    )
    body(
        doc,
        "The technical guide explains the routing rules, which parts of the conversation influence the decision, and how the application handles an unavailable model. This guide does not restate those rules.",
    )
    h2(doc, "Continuing a conversation")
    body(
        doc,
        "You can type another message while a reply is generating and queue it to send afterward. Stop ends the current response. Attachments remain associated with their queued message.",
    )
    body(
        doc,
        "Documents and images supplied in chat become Knowledge for that workspace and can inform subsequent answers. Replies can cite relevant retrieved sources. A general answer does not imply that a source was used when it was not.",
    )
    body(
        doc,
        "Voice transcription supplies text to the composer using the configured service. Review the transcript, and any detected sensitive information, before sending it.",
    )
    figure(doc, "2", "Chat. Status labels sit above the composer. The chevron hides them.")
    h2(doc, "Files from chat")
    body(
        doc,
        "Ask for a file, such as an Excel spreadsheet, a Word document, a CSV or a Markdown note, and the reply carries a file card with a Download button. The browser builds the file when you press Download. Nothing is stored on the server. File names are cleaned, anything that would run when opened is saved as a plain text file, spreadsheet cells that look like formulas are written as text, and a file over 2 MB is refused.",
    )
    body(
        doc,
        "If the model did not write the file format, the browser first tries to build a file from the reply itself, such as a table or a document. Only when that finds nothing, and an administrator has approved a task model, does the helper build it. A request for a file changes the shape of the reply. It does not change what is searched or which tier runs.",
    )
    h2(doc, "A member's first visit")
    body(
        doc,
        "A member's first workspace opens with a short welcome, and a Take the tour button walks through the main areas. The button stops appearing after three visits, or sooner if you close it.",
    )

    h1(doc, "8  Knowledge")
    purpose(doc, "Let conversations draw on documents and images you provide, with relevant sources available to support answers.")
    body(
        doc,
        "Files uploaded through Knowledge or attached in chat belong to the current workspace. The application extracts usable content, applies sensitive-data detection and redaction to that text, and prepares the redacted text for retrieval.",
    )
    body(
        doc,
        "On a local installation the original file is also stored. Extracted text is scrubbed before it is chunked and embedded. Converting an image to a vector does not make the image safe. The extracted content is redacted before embedding. The hosted data path does not store that original through the same local file step.",
    )
    body(
        doc,
        "Uploads are limited before any parser runs: a file is at most 10 MB, and a Word or Excel file is opened with a hard cap on how large it may unpack, so a file that lies about its size is refused. Each person can add a limited number of files a day and store a limited total.",
    )
    body(
        doc,
        "When an answer uses retrieved material, it can identify the supporting source or passage. Source references help users inspect the basis of an answer. They do not guarantee that the answer is correct.",
    )

    h1(doc, "9  Memory")
    purpose(doc, "Show what the application has retained from you, and the context it can use in later conversations.")
    body(
        doc,
        "Memory is a readable view of remembered information and summaries. It makes retained context visible, so later replies are not a guess about what was kept.",
    )
    body(
        doc,
        "Facts come from what you wrote, not from the assistant’s replies. You can correct a fact, pin it, or forget it. Forget hides the fact from later chat and keeps the record so it can be restored. Delete removes the fact. A conversation summary sits alongside the facts.",
    )
    body(
        doc,
        "The vector index used for citations holds document text. It does not hold these fact rows.",
    )

    h1(doc, "10  Bookmarks")
    purpose(doc, "Bookmark useful places in a long conversation and return to them easily.")
    body(
        doc,
        "Bookmark a message when you want to revisit it. Bookmarks lists them. Choosing one returns you to that message in the conversation and highlights it. You can also remove a bookmark.",
    )
    body(
        doc,
        "A bookmark is a navigation aid you curate. It does not become a remembered fact or a Knowledge document.",
    )

    h1(doc, "11  Insights")
    purpose(doc, "Make model activity, estimated costs, and detected sensitive-data findings visible while you work.")
    body(
        doc,
        "Insights is the right-hand panel beside chat, with Models, Usage, Cost, and Security. Each opens a full page for the longer view. It is one of SecureMind’s main transparency features. Collapse it when you want more room for the conversation.",
    )
    body(
        doc,
        "Insights follows the workspace you select, and the date range Today, Weekly, Monthly, or All time. Administrator-wide views are available separately through the Admin console.",
    )
    h2(doc, "Models and the full Usage view")
    body(
        doc,
        "Assign a model to each complexity tier, or choose one Fixed model. Provider chips show how many models of that provider can be assigned. The small blue number is the count of replies routed to that tier for the selected workspace and date range.",
    )
    figure(doc, "3", "Insights, Models. Each tier shows its model and the blue reply count.")
    body(
        doc,
        "Open the full Usage view for charts and details about input tokens, output tokens, requests by model, routing tiers, citation outcomes, and response times.",
    )
    h2(doc, "Cost")
    body(
        doc,
        "Cost provides an approximate breakdown of model API, Bedrock, and voice-transcription usage, using recorded activity and provider rates. Each estimate is raised by 10 percent and rounded up, so it tends high rather than low. It is useful for comparing models, observing spending trends, and judging an approximate spending range. The page also compares what you spent with what the same replies would have cost if every one had used the most expensive default tier. That gap is the routing saving. It is not money returned.",
    )
    body(
        doc,
        "These figures are not billing records. Check the provider’s billing system, or AWS billing for Bedrock, for actual charges. The estimate is not a total cost of ownership for the installation. Local hardware and other infrastructure may cost money outside model usage.",
    )
    body(
        doc,
        "Per-user caps use these estimates as a precaution. When a configured threshold is reached, further billable model calls and voice sessions for that user stop. Ollama calls are not part of that cap. The cap cannot make the provider’s final bill equal the displayed estimate.",
    )
    h2(doc, "Security")
    body(
        doc,
        "Security shows sensitive-data findings for the selected workspace and period. It helps users understand what was detected, where it occurred, and how it was handled. It is not a general vulnerability scanner.",
    )
    body(
        doc,
        "A finding shows the category and what was done with it. It does not show the original sensitive value. A finding can be dismissed or deleted as a log entry. Removing a finding does not restore redacted source content.",
    )

    h1(doc, "12  Model settings")
    purpose(doc, "Manage connections, available models, transcription, and routing in one place.")
    h2(doc, "Connections")
    body(
        doc,
        "Administrators configure vendor APIs and follow the setup guidance for their keys. In the typed-chat path, detected sensitive text is replaced before the model call. This is best-effort protection. Voice audio and some image processing can leave the installation before that text review.",
    )
    body(
        doc,
        "Amazon Bedrock configuration uses the authentication and region described in section 4. Local installations also expose Ollama. Voice transcription settings select Amazon Transcribe or Groq Whisper and explain when audio leaves the device.",
    )
    h2(doc, "Available models")
    body(
        doc,
        "The available-model list shows the models offered through configured connections. Enable, Test, and Check access are three controls, as defined in section 3. Hide and Show manage the chat list. A hidden model can still be the model a tier runs. Check failed means the last check did not succeed.",
    )
    figure(doc, "4", "Settings, Models. Enable, Test, and Check access.")
    body(
        doc,
        "A model a provider has retired is taken off the chat menus automatically the first time a reply fails because of it, and it stays under the not-working group until a Test passes.",
    )
    h2(doc, "Task models")
    body(
        doc,
        "Under Settings, Models, an administrator can approve a helper model for building files, as described in section 4. The setting is off by default. Voice transcription is chosen in the same panel.",
    )
    h2(doc, "Complexity tiers and Fixed model")
    body(
        doc,
        "The larger tier editor provides the same assignment controls as Insights, with more room to browse and filter models. Assign a model to Max, Plus, Standard, Light, and Micro, or select one Fixed model.",
    )

    h1(doc, "13  Detection policy")
    purpose(doc, "Set the sensitivity of future sensitive-data detection across the installation.")
    body(
        doc,
        "The administrator manages this policy under Settings, Detection. Balanced is the default and covers the usual pattern families. High adds passport and national-ID patterns that Balanced does not include. Low narrows detection. For example, Low does not include the street-address and IP patterns that Balanced includes.",
    )
    body(
        doc,
        "Changes apply going forward. Switching from Low to High does not rescan or scrub older messages or documents accepted under Low.",
    )
    body(
        doc,
        "Safe terms let an administrator exempt a known non-sensitive value, such as a domain or a literal the installation should keep. Secrets, Social Security numbers, social insurance numbers, and card numbers cannot be exempted.",
    )
    h2(doc, "How a value is recognized")
    body(
        doc,
        "The review is a pattern scan, not a language model reading the sentence. It looks for shapes that usually mean a personal detail or a secret. When a shape matches and passes that family’s extra check, the span is replaced. A sentence that describes a person without one of these shapes is not caught. The same steps run in the browser before a chat message leaves, and on the server when a file’s text is extracted, before that text is saved. A picture is turned into text first. A voice transcript is reviewed when you send it. The model’s reply is not reviewed.",
    )
    body(
        doc,
        "The steps are fixed. Read the sensitivity. Try only the families that level includes. Search the text for each shape. Drop a hit that fails its extra check. If the family allows it, skip a hit covered by a safe term, and count that skip without storing the value. Merge overlapping hits into one span. Replace the span with a label such as [email redacted]. The label names the category and does not contain the original characters. If nothing remains, the text continues as written. There is no way to restore those characters.",
    )
    add_table(
        doc,
        ["Family", "From", "Extra check", "Safe term?"],
        [
            ["Email", "Low", "Address shape", "Yes"],
            ["US Social Security number", "Low", "3-2-4 with a hyphen or space. No checksum.", "No"],
            ["Canadian social insurance number", "Low", "3-3-3 or nine digits, then the Canadian check digit", "No"],
            ["Card number", "Low", "13 to 19 digits, then the Luhn checksum", "No"],
            ["API keys and private keys", "Low", "A known key form only, not a random string", "No"],
            ["Phone number", "Balanced", "North American shape", "Yes"],
            ["IPv4 address", "Balanced", "Four numbers from 0 to 255", "Yes"],
            ["Street address", "Balanced", "A street number ending in a street word", "Yes"],
            ["Passport or national-ID shape", "High", "One or two letters, then seven to nine digits", "Yes"],
        ],
        [2.0, 0.9, 2.6, 0.9],
    )
    body(
        doc,
        "Luhn is the reason a long number in a log is not treated as a card. The checksum is a weighted sum of the digits, from the right, and it must divide by 10. The Canadian social insurance check is the same style of sum, taken from the left, and it rejects all zeros and all-identical digits. Secrets match a known form: a private-key block, or a known prefix such as sk_live_, ghp_, github_pat_, a Slack xox token, an AWS access-key prefix AKIA or ASIA, a named AWS secret assignment, an OpenAI-style sk- key, or a Google AIza key.",
    )
    body(
        doc,
        "A safe term is an exact phrase, an IPv4 range, a domain, or a prefix. Exact ignores case. A range uses CIDR, such as 10.0.0.0/8. A domain matches that domain or a real boundary, such as .acme.com or @acme.com, and notacme.com does not match acme.com. A prefix matches the start of the text. The install already lists documentation mailboxes (example.com and the reserved endings test, invalid, localhost, and example) and non-routable addresses (private ranges, loopback, link-local, and the documentation ranges). They are visible, and an administrator can remove them. A proposed safe term that itself looks like a Social Security number, a social insurance number, a card, or a secret is refused.",
    )
    figure(doc, "5", "Settings, Detection. Sensitivity and safe terms.")

    h1(doc, "14  Profile and account security")
    purpose(doc, "Manage how you are identified in the application and maintain your account details.")
    body(
        doc,
        "Profile lets you set a profile image, change your display name, and add a description about yourself. The photo is available to a signed-in session. It is not a public link. Account security, where the installation uses its sign-in service, lets you change your password and turn on two-step verification with an authenticator app.",
    )
    body(
        doc,
        "Two-step verification is optional and uses real Amazon Cognito, so it is available on the hosted edition. The local Docker sign-in does not support codes, and the page hides the control there. You set it up by scanning a QR code and typing a code to confirm. At sign-in, after your password, you type the six-digit code. Turning it off needs your password again, and repeated wrong passwords are slowed down. Administrators see a recommendation until they turn it on, and the Owner can reset someone else’s second step from the Admin page if a phone is lost.",
    )

    h1(doc, "15  Appearance")
    purpose(doc, "Choose a comfortable display theme.")
    body(
        doc,
        "Open Settings and choose Appearance to switch between light and dark. The choice is saved on the device. It does not change other users’ screens. Dark is the default until a choice is saved.",
    )

    h1(doc, "16  Admin console")
    purpose(doc, "Manage the people using the installation and oversee their usage and estimated spending.")
    body(
        doc,
        "Administrators can view users’ email addresses, monthly estimated costs, and Insights for the installation or for one person. They can set a per-user cost cap.",
    )
    body(
        doc,
        "The administrator creates accounts, assigns roles, sets passwords, and can suspend an account. Members use the shared connections and cannot view or change provider secret keys.",
    )
    body(
        doc,
        "Only the Owner can change another administrator’s role, status, password, or two-step verification, so one administrator cannot take over another. A suspended person is signed out within seconds. Every administrative change is kept in an audit log. Setting a member’s password lets an administrator sign in as that member, so keep the number of administrators small and have them turn on two-step verification.",
    )
    body(
        doc,
        "Cost caps are precautionary limits based on estimated use. The application blocks covered billable requests when the estimate reaches the cap. Administrators should still review actual charges with the provider. These limits are not provider-enforced billing budgets.",
    )
    body(
        doc,
        "Add user creates an account with a password the administrator types or generates, and sends a welcome email when outgoing email is set up. A deleted person can be restored. A person who is deleted loses sign-in, but their workspaces, Knowledge and history are not erased.",
    )
    body(
        doc,
        "The Portal card sets the logo, name, and short phrase shown in the top-left of the workspace. The sign-in page, the browser title, and email stay SecureMind.",
    )
    figure(doc, "6", "Admin. Directory, roles, and estimated caps.")

    h1(doc, "17  Getting started")
    purpose(doc, "Help a new user reach a working model and understand how to begin.")
    body(
        doc,
        "Getting started appears while chat is not ready. It points toward Settings, Setup, and Models: connect a model source, make a model available, and send a first message. Sensitive-data review and Insights are part of that introduction. A member who joins an installation that is already set up sees a short welcome and can take a tour. The full setup steps are on the website, as described in section 5.",
    )

    h1(doc, "18  How the parts fit")
    body(
        doc,
        "Start with your own installation. The administrator connects services, configures detection, creates accounts, and sets estimated spending limits. Members then work in a workspace using that configuration.",
    )
    body(
        doc,
        "Chat reviews typed text before processing. Knowledge supplies uploaded content. Memory retains what you wrote. Bookmarks return you to a message. A complexity tier or a fixed model determines which model answers. Insights shows the activity, the estimated cost, and the findings.",
    )
    h2(doc, "Installation instructions")
    body(
        doc,
        "Setup covers a Docker installation on your own machine and a deployment in your own AWS account. Those steps are on the SecureMind website and linked from the GitHub README. The cloud region is set in the infrastructure configuration, as described in section 5. The technical guide describes provider regions and processing destinations.",
    )

    h1(doc, "19  Scope and limitations")
    body(
        doc,
        "SecureMind is an additional layer of control, reassurance, and transparency for people using AI in their own environment. It is not a compliance program, a certification, or a replacement for an organization’s required security controls.",
    )
    body(
        doc,
        "Detection is best effort. A value that is not detected may be processed or stored. Browser review is the path inside the application. For a request that bypasses it, the server catches only the highest-risk values, not every category.",
    )
    body(
        doc,
        "Costs and caps rely on estimates, not invoices. Actual billing comes from the provider. A locally hosted application may still send content to external services when those services are configured.",
    )
    body(
        doc,
        "Each installation is independently operated. SecureMind is not intended to run a shared public service that holds every customer’s API keys, and one installation is not a boundary between separate organizations. A region set for the deployment does not, by itself, mean that every connected service processes data only there. An AWS web application firewall is not included; an administrator can add one to their own account.",
    )
    body(
        doc,
        "SecureMind does not accept pasted AWS access keys. Local Bedrock uses the credential chain in section 4. Raw audio sent for hosted transcription is processed before the transcript can be reviewed and redacted. The model’s reply is not scanned for sensitive values, and a signed-in session stays valid until it expires, up to an hour, even after you sign out.",
    )

    try:
        doc.save(OUT)
        print(OUT)
    except PermissionError:
        fallback = OUT.with_name("SecureMind-Feature-Guide-v1.docx")
        doc.save(fallback)
        print(fallback)


if __name__ == "__main__":
    build()
